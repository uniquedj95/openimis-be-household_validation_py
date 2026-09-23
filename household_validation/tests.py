import base64
from contextlib import nullcontext
from dataclasses import replace
from datetime import date, datetime, timezone as datetime_timezone
from io import BytesIO
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import MagicMock, call, patch
from uuid import uuid4

from django.db.models import Q
from django.test import TestCase as DjangoTestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.models import User
from individual.models import Group, GroupIndividual, Individual
from location.models import Hotspot, HotspotVillage, Location, MicroCatchment

from household_validation.apps import (
    DEFAULT_CONFIG,
    DISTRICT_VALIDATION_ROLE_RIGHTS,
    DISTRICT_VALIDATION_ROLES,
    GROUP_RIGHTS,
    HOUSEHOLD_VALIDATION_RIGHTS,
    ROLE_DISTRICT_ADMINISTRATOR,
    ROLE_DISTRICT_PROGRAM_MANAGER,
    ROLE_DISTRICT_USER,
    RIGHT_GROUP_SEARCH,
    RIGHT_GROUP_UPDATE,
    RIGHT_HOUSEHOLD_VALIDATION_ERROR_REPORT,
    RIGHT_HOUSEHOLD_VALIDATION_HISTORY,
    RIGHT_HOUSEHOLD_VALIDATION_QUERY_EXPORT,
    RIGHT_HOUSEHOLD_VALIDATION_UPLOAD,
    HouseholdValidationConfig,
)
from household_validation.excel import (
    BUSINESS_COLUMNS,
    BUSINESS_DURATION_COLUMN,
    BUSINESS_TYPE_COLUMN,
    BUSINESS_TYPE_OPTIONS_SHEET,
    HAS_BUSINESS_COLUMN,
    EXCEL_COLUMNS,
    HOUSEHOLD_ROW_COLORS,
    HOTSPOT_COLUMN,
    MICRO_CATCHMENT_COLUMN,
    PROJECT_OPTIONS_SHEET,
    ExcelValidationListExporter,
    build_rejected_households_workbook_bytes,
)
from household_validation.gql_permissions import (
    HouseholdValidationPermissionError,
    require_permissions,
)
from household_validation.gql_mutations import GenerateHouseholdValidationListMutation
from household_validation.models import HouseholdValidationBatchRow
from household_validation.project_lookup import (
    ACTIVE_PROJECT_STATUSES,
    ProjectOption,
    project_option_from_project,
)
from household_validation.selection import (
    CATEGORY_FEMALE_HEADED,
    CATEGORY_OTHER,
    CATEGORY_YOUTH,
    ROW_TYPE_MAIN,
    ROW_TYPE_RESERVE,
    EligibleHousehold,
    EligibleMember,
    SelectedHousehold,
    SelectionResult,
    is_truthy,
    select_households,
)
from household_validation.services import (
    EligibleHouseholdSelectionService,
    HouseholdValidationUploadService,
    _json_safe,
    _local_date,
    build_validation_error_report,
)
from household_validation.schema import Query
from household_validation.upload import (
    OPTIONAL_UPLOAD_COLUMNS,
    PROJECT_SELECTION_TYPE_INTENT,
    VALIDATION_STATUS_NOT_VERIFIED,
    VALIDATION_LIST_SHEET,
    UploadedValidationRow,
    build_validation_error_report_csv,
    build_validation_json_ext,
    member_structural_errors,
    parse_validation_workbook,
)
from household_validation.wealth import get_household_pmt_score, get_household_wealth_quintile


class HouseholdValidationConfigTest(TestCase):
    def test_business_columns_are_disabled_by_default_and_for_older_config(self):
        self.assertIs(DEFAULT_CONFIG["business_columns_enabled"], False)
        with patch.object(HouseholdValidationConfig, "business_columns_enabled", True):
            HouseholdValidationConfig._load_config({})
            self.assertIs(HouseholdValidationConfig.business_columns_enabled, False)

    def test_business_columns_can_be_enabled_per_deployment(self):
        with patch.object(HouseholdValidationConfig, "business_columns_enabled", False):
            HouseholdValidationConfig._load_config({"business_columns_enabled": True})
            self.assertIs(HouseholdValidationConfig.business_columns_enabled, True)

    def test_default_config_exposes_permission_lists(self):
        self.assertEqual(
            DEFAULT_CONFIG["gql_query_household_validation_rule_perms"],
            [str(RIGHT_HOUSEHOLD_VALIDATION_QUERY_EXPORT)],
        )
        self.assertEqual(
            DEFAULT_CONFIG["gql_mutation_generate_household_validation_list_perms"],
            [str(RIGHT_HOUSEHOLD_VALIDATION_QUERY_EXPORT)],
        )
        self.assertEqual(
            DEFAULT_CONFIG["gql_mutation_upload_household_validation_list_perms"],
            [str(RIGHT_HOUSEHOLD_VALIDATION_UPLOAD)],
        )
        self.assertEqual(
            DEFAULT_CONFIG["gql_query_household_validation_history_perms"],
            [str(RIGHT_HOUSEHOLD_VALIDATION_HISTORY)],
        )
        self.assertEqual(
            DEFAULT_CONFIG["gql_query_household_validation_error_report_perms"],
            [str(RIGHT_HOUSEHOLD_VALIDATION_ERROR_REPORT)],
        )
        self.assertEqual(
            DEFAULT_CONFIG["group_search_perms"],
            [str(RIGHT_GROUP_SEARCH)],
        )
        self.assertEqual(
            DEFAULT_CONFIG["group_update_perms"],
            [str(RIGHT_GROUP_UPDATE)],
        )

    def test_default_config_exposes_selection_percentages(self):
        quota_config = DEFAULT_CONFIG["program_eligibility_rules"]["PWP"]["selection_strategy"]
        self.assertEqual(quota_config["female_headed_percentage"], 40)
        self.assertEqual(quota_config["youth_headed_percentage"], 40)
        self.assertEqual(quota_config["reserve_percentage"], 20)

    def test_right_sets_match_default_config_scope(self):
        self.assertEqual(
            HOUSEHOLD_VALIDATION_RIGHTS,
            [
                RIGHT_HOUSEHOLD_VALIDATION_QUERY_EXPORT,
                RIGHT_HOUSEHOLD_VALIDATION_UPLOAD,
                RIGHT_HOUSEHOLD_VALIDATION_HISTORY,
                RIGHT_HOUSEHOLD_VALIDATION_ERROR_REPORT,
            ],
        )
        self.assertIn(RIGHT_GROUP_SEARCH, GROUP_RIGHTS)
        self.assertIn(RIGHT_GROUP_UPDATE, GROUP_RIGHTS)

    def test_district_validation_roles_have_validation_and_group_update_rights(self):
        self.assertEqual(
            DISTRICT_VALIDATION_ROLES,
            [
                ROLE_DISTRICT_ADMINISTRATOR,
                ROLE_DISTRICT_PROGRAM_MANAGER,
                ROLE_DISTRICT_USER,
            ],
        )
        for role_name in DISTRICT_VALIDATION_ROLES:
            role_rights = DISTRICT_VALIDATION_ROLE_RIGHTS[role_name]
            for right in HOUSEHOLD_VALIDATION_RIGHTS:
                self.assertIn(right, role_rights)
            self.assertIn(RIGHT_GROUP_SEARCH, role_rights)
            self.assertIn(RIGHT_GROUP_UPDATE, role_rights)


class RejectedBatchRowsQueryTest(TestCase):
    @patch("household_validation.schema.HouseholdValidationBatchRow.objects.filter")
    def test_rejected_rows_use_upload_permission_without_changing_history_query(
        self,
        filter_mock,
    ):
        user = MagicMock()
        info = MagicMock()
        info.context.user = user
        rejected_row = MagicMock()
        rejected_row.json_ext = {"error_code": "MULTIPLE_PRIMARY_WORKERS"}
        rejected_row.raw_row = {
            "form_number": "FORM-001",
            "group_uuid": "group-1",
        }
        rejected_row.group_id = "group-1"
        rejected_row.row_number = 2
        rejected_row.error_message = (
            "household has more than one primary worker"
        )
        upload_attempt_id = "01a03ab4-c5ac-7b78-a485-f321e7d092f8"
        filter_mock.return_value.order_by.return_value = [rejected_row]

        with patch.object(Query, "_check_permissions") as check_permissions_mock:
            result = Query.resolve_household_validation_rejected_batch_rows(
                None,
                info,
                batch_id="01a03aae-4896-7f00-9357-78f69fb8e6ca",
                upload_attempt_id=upload_attempt_id,
            )

        check_permissions_mock.assert_called_once_with(
            user,
            HouseholdValidationConfig.gql_mutation_upload_household_validation_list_perms,
        )
        filter_mock.assert_called_once_with(
            batch_id="01a03aae-4896-7f00-9357-78f69fb8e6ca",
            upload_attempt_id=upload_attempt_id,
            status=HouseholdValidationBatchRow.Status.REJECTED,
            is_deleted=False,
        )
        filter_mock.return_value.order_by.assert_called_once_with("row_number")
        self.assertEqual(result.rows, [rejected_row])
        self.assertEqual(result.count, 1)
        self.assertEqual(
            result.file_name,
            (
                "rejected_households_01a03aae-4896-7f00-9357-78f69fb8e6ca_"
                "01a03ab4-c5ac-7b78-a485-f321e7d092f8.xlsx"
            ),
        )
        workbook = load_workbook(BytesIO(base64.b64decode(result.file_base64)))
        self.assertEqual(workbook.active.title, "Rejected Households")
        self.assertEqual(workbook.active.max_row, 2)


class RejectedHouseholdsWorkbookTest(TestCase):
    def test_contains_only_multiple_primary_worker_rejections(self):
        rejection_message = (
            "household has more than one primary worker"
        )
        rows = [
            SimpleNamespace(
                row_number=4,
                group_id="group-1",
                raw_row={"form_number": "FORM-001", "member_name": "Member One"},
                json_ext={"error_code": "MULTIPLE_PRIMARY_WORKERS"},
                error_message=rejection_message,
            ),
            SimpleNamespace(
                row_number=5,
                group_id="group-1",
                raw_row={"form_number": "FORM-001", "member_name": "Member Two"},
                json_ext={"error_code": "MULTIPLE_PRIMARY_WORKERS"},
                error_message=rejection_message,
            ),
            SimpleNamespace(
                row_number=6,
                group_id="group-2",
                raw_row={"form_number": "FORM-002", "member_name": "Other Error"},
                json_ext={"error_code": "MEMBER_NOT_FOUND"},
                error_message="Member was not found",
            ),
        ]

        report, rejected_row_count = build_rejected_households_workbook_bytes(rows)
        workbook = load_workbook(BytesIO(report))
        worksheet = workbook["Rejected Households"]
        values = list(worksheet.values)

        self.assertEqual(rejected_row_count, 1)
        self.assertEqual(worksheet.max_row, 2)
        self.assertIn("rejection_reason", values[0])
        self.assertEqual(values[1][0], "FORM-001")
        self.assertEqual(values[1][1], "group-1")
        self.assertEqual(values[1][2], "4, 5")
        self.assertNotIn("FORM-002", str(values))


def _dob_for_age(age):
    today = date.today()
    return date(today.year - age, today.month, today.day)


def _member(
    member_id,
    gender="Male",
    age=40,
    fit_for_work=True,
    role=None,
    recipient_type=None,
    data_source=None,
    json_ext=None,
):
    return EligibleMember(
        id=member_id,
        gender=gender,
        dob=_dob_for_age(age),
        fit_for_work=fit_for_work,
        role=role,
        recipient_type=recipient_type,
        data_source=data_source,
        json_ext=json_ext if json_ext is not None else {},
    )


def _household(
    household_id,
    wealth_quintile,
    head_gender="Male",
    eligible_member_age=40,
    last_verified_date=None,
    eligible_members=None,
    village_id=None,
    village_code=None,
    village_name=None,
):
    head = _member(
        f"{household_id}-head",
        gender=head_gender,
        age=50,
        role="HEAD",
    )
    if eligible_members is None:
        eligible_members = [
            _member(
                f"{household_id}-worker",
                gender="Male",
                age=eligible_member_age,
            )
        ]
    return EligibleHousehold(
        id=household_id,
        code=str(household_id),
        wealth_quintile=wealth_quintile,
        last_verified_date=last_verified_date,
        head=head,
        eligible_members=eligible_members,
        village_id=village_id,
        village_code=village_code,
        village_name=village_name,
    )


class HouseholdSelectionTest(TestCase):
    def test_select_households_allocates_target_proportionally_by_village(self):
        households = (
            [
                _household(f"a-{index}", "Poorest", village_code="A", village_name="Village A")
                for index in range(30)
            ]
            + [
                _household(f"b-{index}", "Poorest", village_code="B", village_name="Village B")
                for index in range(15)
            ]
            + [
                _household(f"c-{index}", "Poorest", village_code="C", village_name="Village C")
                for index in range(10)
            ]
        )

        result, summary = select_households(
            households,
            target_count=20,
            allocate_by_village=True,
            rule={"selection_strategy": {}},
        )

        selected_by_village = {
            code: len([
                row
                for row in result.main
                if row.household.village_code == code
            ])
            for code in ("A", "B", "C")
        }
        self.assertEqual(selected_by_village, {"A": 11, "B": 5, "C": 4})
        self.assertEqual(summary["selected_households"], 20)
        self.assertEqual(
            {
                row["village_code"]: row["allocated_households"]
                for row in summary["village_breakdown"]
            },
            {"A": 11, "B": 5, "C": 4},
        )

    def test_village_allocation_guarantees_one_household_when_target_allows(self):
        households = (
            [
                _household(f"a-{index}", "Poorest", village_code="A")
                for index in range(98)
            ]
            + [_household("b-1", "Poorest", village_code="B")]
            + [_household("c-1", "Poorest", village_code="C")]
        )

        result, _ = select_households(
            households,
            target_count=3,
            allocate_by_village=True,
            rule={"selection_strategy": {}},
        )

        self.assertEqual(
            {row.household.village_code for row in result.main},
            {"A", "B", "C"},
        )

    def test_allocate_by_village_config_can_disable_village_allocation(self):
        households = (
            [
                _household(f"a-{index}", "Poorest", village_code="A")
                for index in range(98)
            ]
            + [_household("b-1", "Poorest", village_code="B")]
            + [_household("c-1", "Poorest", village_code="C")]
        )

        result, _ = select_households(
            households,
            target_count=3,
            allocate_by_village=True,
            rule={"selection_strategy": {"allocate_by_village": False}},
        )

        # With village allocation disabled via config, the request-level
        # allocate_by_village=True is overridden: selection falls back to
        # the flat, wealth-ranked quota algorithm, so village B/C no longer
        # get a guaranteed slot.
        self.assertEqual(
            {row.household.village_code for row in result.main},
            {"A"},
        )

    def test_village_reserve_is_proportional_and_excludes_main_households(self):
        households = [
            _household(
                f"{code}-{index}",
                "Poorest",
                village_code=code,
            )
            for code in ("A", "B")
            for index in range(10)
        ]

        result, summary = select_households(
            households,
            target_count=10,
            allocate_by_village=True,
            rule={"selection_strategy": {"reserve_percentage": 20}},
        )

        self.assertEqual(
            {
                code: len([
                    row
                    for row in result.main
                    if row.household.village_code == code
                ])
                for code in ("A", "B")
            },
            {"A": 5, "B": 5},
        )
        self.assertEqual(
            {
                code: len([
                    row
                    for row in result.reserve
                    if row.household.village_code == code
                ])
                for code in ("A", "B")
            },
            {"A": 1, "B": 1},
        )
        self.assertFalse(
            {row.household.id for row in result.main}
            & {row.household.id for row in result.reserve}
        )
        self.assertEqual(summary["reserve_households"], 2)

    def test_select_households_applies_quota_categories(self):
        households = [
            _household("female", "Poorest", head_gender="Female"),
            _household("youth", "Poorest", eligible_member_age=25),
            _household("other", "Poorest"),
        ]

        result, _ = select_households(
            households,
            target_count=3,
            rule={"selection_strategy": {}},
        )

        self.assertEqual([row.row_type for row in result.main], [ROW_TYPE_MAIN] * 3)
        self.assertEqual(
            [row.category for row in result.main],
            [CATEGORY_FEMALE_HEADED, CATEGORY_YOUTH, CATEGORY_OTHER],
        )

    def test_select_households_prioritizes_poorest_within_category(self):
        households = [
            _household("female-richer", "Richer", head_gender="Female"),
            _household("female-poorest", "Poorest", head_gender="Female"),
            _household("youth", "Poorest", eligible_member_age=20),
        ]

        result, _ = select_households(
            households,
            target_count=1,
            rule={"selection_strategy": {}},
        )

        self.assertEqual(result.main[0].household.id, "female-poorest")

    def test_select_households_defaults_target_to_all_eligible_and_caps_reserve(self):
        households = [
            _household("female", "Poorest", head_gender="Female"),
            _household("youth", "Poorer", eligible_member_age=20),
        ]

        result, _ = select_households(households, rule={"selection_strategy": {}})

        self.assertEqual(len(result.main), 2)
        self.assertEqual(result.reserve, [])

    def test_select_households_adds_reserve_from_remaining_households(self):
        households = [
            _household("female", "Poorest", head_gender="Female"),
            _household("youth", "Poorer", eligible_member_age=20),
            _household("other", "Middle"),
        ]

        result, summary = select_households(
            households,
            target_count=2,
            rule={"selection_strategy": {"reserve_percentage": 50}},
        )

        self.assertEqual(len(result.main), 2)
        self.assertEqual(len(result.reserve), 1)
        self.assertEqual(result.reserve[0].row_type, ROW_TYPE_RESERVE)
        self.assertEqual(result.reserve[0].household.id, "other")
        self.assertEqual(summary["selected_households"], 2)
        self.assertEqual(summary["selected_individuals"], 2)
        self.assertEqual(summary["reserve_households"], 1)

    def test_select_households_adds_default_twenty_percent_reserve(self):
        households = [
            _household(index, "Poorest")
            for index in range(1, 21)
        ]

        result, _ = select_households(
            households,
            target_count=10,
            rule={"selection_strategy": {}},
        )

        self.assertEqual(len(result.main), 10)
        self.assertEqual(len(result.reserve), 2)
        self.assertEqual(result.reserve[0].row_type, ROW_TYPE_RESERVE)

    def test_reserve_capacity_uses_actual_main_selection_count(self):
        households = [
            _household(f"household-{index}", "Poorest")
            for index in range(3)
        ]
        selected_main = [
            SelectedHousehold(
                household=household,
                category=CATEGORY_OTHER,
                row_type=ROW_TYPE_MAIN,
            )
            for household in households[:2]
        ]
        category_counts = {
            CATEGORY_FEMALE_HEADED: 0,
            CATEGORY_YOUTH: 0,
            CATEGORY_OTHER: 2,
        }
        household_categories = {
            household.id: CATEGORY_OTHER for household in households
        }

        with patch(
            "household_validation.selection._select_main_households",
            return_value=(
                selected_main,
                {household.id for household in households[:2]},
                category_counts,
                sum(
                    len(household.eligible_members)
                    for household in households[:2]
                ),
                household_categories,
            ),
        ):
            result, summary = select_households(
                households,
                target_count=3,
                rule={"selection_strategy": {"reserve_percentage": 100}},
            )

        self.assertEqual(
            [row.household.id for row in result.reserve],
            [households[2].id],
        )
        self.assertEqual(summary["reserve_households"], 1)

    def test_select_households_excludes_recently_verified_households(self):
        households = [
            _household(
                "recent",
                "Poorest",
                last_verified_date=date(2026, 7, 2),
            ),
            _household(
                "old",
                "Poorer",
                last_verified_date=date(2026, 6, 1),
            ),
        ]

        result, _ = select_households(
            households,
            exclude_verified_after=date(2026, 6, 30),
        )

        self.assertEqual([row.household.id for row in result.main], ["old"])

    def test_select_households_ignores_households_without_eligible_members(self):
        households = [
            EligibleHousehold(
                id="no-workers",
                code="no-workers",
                wealth_quintile="Poorest",
                eligible_members=[],
            ),
            _household("worker", "Poorer"),
        ]

        result, _ = select_households(households)

        self.assertEqual([row.household.id for row in result.main], ["worker"])

    def test_select_households_backfills_quota_shortage_from_other_categories(self):
        households = [
            _household("female", "Poorest", head_gender="Female"),
            _household("other-1", "Poorer"),
            _household("other-2", "Middle"),
        ]

        result, summary = select_households(
            households,
            target_count=3,
            rule={"selection_strategy": {}},
        )

        self.assertEqual(len(result.main), 3)
        self.assertEqual(
            [row.household.id for row in result.main],
            ["female", "other-1", "other-2"],
        )
        # "other-2" is only picked up by the backfill loop (there's no youth
        # household to fill that quota slot), so the summary counts must
        # reflect it too, not just the households chosen during quota-fill.
        self.assertEqual(summary["selected_households"], 3)
        self.assertEqual(summary["selected_female_headed_households"], 1)
        self.assertEqual(summary["selected_youth_households"], 0)
        self.assertEqual(summary["selected_other_households"], 2)

    def test_selection_result_expands_selected_households_to_member_rows(self):
        households = [
            _household(
                "household",
                "Poorest",
                eligible_members=[
                    _member("worker-1", age=24),
                    _member("worker-2", age=45),
                ],
            )
        ]

        result, _ = select_households(households)

        self.assertEqual(len(result.member_rows), 2)
        self.assertEqual(
            [row.member.id for row in result.member_rows],
            ["worker-1", "worker-2"],
        )
        self.assertEqual(
            {row.household.id for row in result.member_rows},
            {"household"},
        )

    def test_select_households_treats_ages_18_to_35_as_youth(self):
        households = [
            _household("age-18", "Poorest", eligible_member_age=18),
            _household("age-35", "Poorest", eligible_member_age=35),
            _household("age-36", "Poorest", eligible_member_age=36),
        ]

        result, _ = select_households(
            households,
            target_count=3,
            rule={"selection_strategy": {}},
        )

        categories = {
            row.household.id: row.category
            for row in result.main
        }
        self.assertEqual(categories["age-18"], CATEGORY_YOUTH)
        self.assertEqual(categories["age-35"], CATEGORY_YOUTH)
        self.assertEqual(categories["age-36"], CATEGORY_OTHER)

    def test_is_truthy_accepts_ubr_boolean_shapes(self):
        for value in (True, 1, "1", "true", "TRUE", "yes", "Y"):
            self.assertTrue(is_truthy(value))
        for value in (False, 0, "0", "false", "no", None, ""):
            self.assertFalse(is_truthy(value))

    def test_select_households_uses_configured_percentages(self):
        households = [
            _household("female", "Poorest", head_gender="Female"),
            _household("youth", "Poorest", eligible_member_age=25),
            _household("other", "Poorest"),
        ]

        result, _ = select_households(
            households,
            target_count=2,
            rule={"selection_strategy": {"female_headed_percentage": 0, "youth_headed_percentage": 100}},
        )

        self.assertEqual(
            [row.category for row in result.main],
            [CATEGORY_YOUTH, CATEGORY_FEMALE_HEADED],
        )

    def test_other_quota_absorbs_the_remainder_when_configured_percentages_total_below_100(self):
        households = (
            [_household(f"female-{i}", "Poorest", head_gender="Female") for i in range(5)]
            + [_household(f"youth-{i}", "Poorest", eligible_member_age=25) for i in range(5)]
            + [_household(f"other-{i}", "Poorest") for i in range(5)]
        )

        result, summary = select_households(
            households,
            target_count=10,
            rule={"selection_strategy": {"female_headed_percentage": 30, "youth_headed_percentage": 30}},
        )

        self.assertEqual(len(result.main), 10)
        self.assertEqual(summary["selected_female_headed_households"], 3)
        self.assertEqual(summary["selected_youth_households"], 3)
        self.assertEqual(summary["selected_other_households"], 4)

    def test_female_and_youth_quotas_are_normalized_when_configured_percentages_exceed_100(self):
        households = (
            [_household(f"female-{i}", "Poorest", head_gender="Female") for i in range(10)]
            + [_household(f"youth-{i}", "Poorest", eligible_member_age=25) for i in range(10)]
        )

        result, summary = select_households(
            households,
            target_count=10,
            rule={"selection_strategy": {"female_headed_percentage": 70, "youth_headed_percentage": 60}},
        )

        self.assertEqual(len(result.main), 10)
        self.assertEqual(summary["selected_other_households"], 0)
        self.assertEqual(
            summary["selected_female_headed_households"] + summary["selected_youth_households"],
            10,
        )


class EligibleMemberIsEligibleTest(TestCase):
    def test_no_rule_falls_back_to_fit_for_work(self):
        self.assertTrue(_member("m1", fit_for_work=True).is_eligible(None))
        self.assertFalse(_member("m1", fit_for_work=False).is_eligible(None))
        self.assertTrue(_member("m1", fit_for_work=True).is_eligible({}))

    def test_requires_data_source_matches_case_insensitively(self):
        member = _member("m1", data_source="pwp")
        self.assertTrue(member.is_eligible({"requires_data_source": "PWP"}))
        self.assertFalse(member.is_eligible({"requires_data_source": "SCTP"}))

    def test_requires_data_source_fails_when_member_has_none(self):
        member = _member("m1", data_source=None)
        self.assertFalse(member.is_eligible({"requires_data_source": "PWP"}))

    def test_requires_recipient_type_matches_case_insensitively(self):
        member = _member("m1", recipient_type="primary")
        self.assertTrue(member.is_eligible({"requires_recipient_type": "PRIMARY"}))
        self.assertFalse(member.is_eligible({"requires_recipient_type": "SECONDARY"}))

    def test_member_flag_requires_truthy_json_ext_value(self):
        rule = {"member_flag": "business_experience"}
        eligible = _member("m1", json_ext={"business_experience": True})
        ineligible = _member("m2", json_ext={"business_experience": False})
        missing = _member("m3", json_ext={})
        self.assertTrue(eligible.is_eligible(rule))
        self.assertFalse(ineligible.is_eligible(rule))
        self.assertFalse(missing.is_eligible(rule))

    def test_age_bounds_are_inclusive(self):
        rule = {"member_min_age": 18, "member_max_age": 60}
        self.assertTrue(_member("m1", age=18).is_eligible(rule))
        self.assertTrue(_member("m2", age=60).is_eligible(rule))
        self.assertFalse(_member("m3", age=17).is_eligible(rule))
        self.assertFalse(_member("m4", age=61).is_eligible(rule))

    def test_age_bounds_exclude_members_with_no_dob(self):
        member = replace(_member("m1", age=30), dob=None)
        self.assertFalse(member.is_eligible({"member_min_age": 18}))

    def test_rmep_like_rule_requires_data_source_but_not_business_experience(self):
        rule = {"requires_data_source": "PWP", "priority_flag": "business_experience"}
        without_business = _member("m1", data_source="PWP", json_ext={"business_experience": False})
        self.assertTrue(without_business.is_eligible(rule))
        wrong_source = _member("m2", data_source="SCTP", json_ext={"business_experience": True})
        self.assertFalse(wrong_source.is_eligible(rule))

    def test_upg_like_rule_requires_data_source_fit_for_work_and_age_range(self):
        rule = {
            "requires_data_source": "SCTP",
            "member_flag": "fit_for_work",
            "member_min_age": 18,
            "member_max_age": 60,
        }
        eligible = _member("m1", age=30, data_source="SCTP", json_ext={"fit_for_work": True})
        self.assertTrue(eligible.is_eligible(rule))
        too_young = _member("m2", age=17, data_source="SCTP", json_ext={"fit_for_work": True})
        self.assertFalse(too_young.is_eligible(rule))
        not_fit = _member("m3", age=30, data_source="SCTP", json_ext={"fit_for_work": False})
        self.assertFalse(not_fit.is_eligible(rule))


class ResolveEligibilityRuleTest(TestCase):
    """Covers EligibleHouseholdSelectionService._resolve_eligibility_rule against
    the module's real default program_eligibility_rules config, so it also acts
    as a regression check on the shipped PWP/RMEP/UPG defaults themselves."""

    def test_no_benefit_plan_code_resolves_to_pwp(self):
        service = EligibleHouseholdSelectionService()
        rule = service._resolve_eligibility_rule(None)
        self.assertEqual(rule, DEFAULT_CONFIG["program_eligibility_rules"]["PWP"])
        self.assertIn("selection_strategy", rule)

    def test_unmatched_benefit_plan_code_falls_back_to_pwp(self):
        service = EligibleHouseholdSelectionService()
        rule = service._resolve_eligibility_rule("SOME_OTHER_PROGRAM")
        self.assertEqual(rule, DEFAULT_CONFIG["program_eligibility_rules"]["PWP"])

    def test_matched_benefit_plan_code_is_case_insensitive(self):
        service = EligibleHouseholdSelectionService()
        expected = DEFAULT_CONFIG["program_eligibility_rules"]["RMEP"]
        self.assertEqual(service._resolve_eligibility_rule("RMEP"), expected)
        self.assertEqual(service._resolve_eligibility_rule("rmep"), expected)
        self.assertEqual(service._resolve_eligibility_rule("Rmep"), expected)

    def test_upg_resolves_to_its_own_rule(self):
        service = EligibleHouseholdSelectionService()
        self.assertEqual(
            service._resolve_eligibility_rule("UPG"),
            DEFAULT_CONFIG["program_eligibility_rules"]["UPG"],
        )

    def test_missing_program_eligibility_rules_config_resolves_to_empty_rule(self):
        service = EligibleHouseholdSelectionService()
        with patch.object(HouseholdValidationConfig, "program_eligibility_rules", None):
            self.assertEqual(service._resolve_eligibility_rule(None), {})
            self.assertEqual(service._resolve_eligibility_rule("RMEP"), {})


class HouseholdValidationPreviewServiceTest(TestCase):
    def test_generate_and_preview_share_selection_result(self):
        service = _FakeSelectionService(
            [
                _fake_group("group-1", "HH-001", "Female", 30, "Poorest"),
                _fake_group("group-2", "HH-002", "Male", 22, "Poorer"),
                _fake_group("group-3", "HH-003", "Male", 45, "Middle"),
            ]
        )

        filters = {
            "target_count": 2,
        }
        _, summary = service.generate(**filters)
        preview_rows = service.preview(**filters)

        self.assertEqual(summary["total_households"], 3)
        self.assertEqual(summary["total_individuals"], 3)
        self.assertEqual(summary["eligible_households"], 3)
        self.assertEqual(summary["selected_households"], 2)
        self.assertEqual(summary["selected_individuals"], 2)
        self.assertEqual(summary["reserve_households"], 1)
        self.assertEqual(len(preview_rows), 3)
        self.assertEqual(
            len([row for row in preview_rows if row.row_type == ROW_TYPE_MAIN]),
            summary["selected_individuals"],
        )

    def test_generate_totals_cover_full_catchment_before_optional_filters(self):
        catchment_groups = [
            _fake_group("group-1", "HH-001", "Female", 30, "Poorest"),
            _fake_group("group-2", "HH-002", "Male", 22, "Poorer"),
        ]
        service = _FakeSelectionService(catchment_groups)
        filtered_queryset = _FakeQuerySet([catchment_groups[0]])

        with (
            patch.object(
                service,
                "_apply_micro_catchment_scope",
                return_value=_FakeQuerySet(catchment_groups),
            ) as catchment_scope_mock,
            patch.object(
                service,
                "_apply_location_filters",
                return_value=filtered_queryset,
            ),
        ):
            _, summary = service.generate(
                catchment_code="MC01",
                village_codes=["V01"],
                target_count=1,
            )

        catchment_scope_mock.assert_called_once()
        self.assertEqual(summary["total_households"], 2)
        self.assertEqual(summary["total_individuals"], 2)
        self.assertEqual(summary["eligible_households"], 1)
        self.assertEqual(summary["selected_households"], 1)

    def test_preview_rows_include_location_and_validation_fields(self):
        service = _FakeSelectionService(
            [
                _fake_group(
                    "group-1",
                    "HH-001",
                    "Female",
                    30,
                    "Poorest",
                    validation_status="VERIFIED",
                )
            ]
        )

        row = service.preview(target_count=1)[0]

        self.assertEqual(row.group_uuid, "group-1")
        self.assertEqual(row.group_code, "HH-001")
        self.assertEqual(row.head_name, "Head Person")
        self.assertEqual(row.individual_first_name, "Head")
        self.assertEqual(row.region, "Region")
        self.assertEqual(row.district, "District")
        self.assertEqual(row.municipality, "Traditional Authority")
        self.assertEqual(row.village, "Village")
        self.assertEqual(row.wealth_quintile, "Poorest")
        self.assertEqual(row.validation_status, "VERIFIED")


class ProgramBasedGenerationIntegrationTest(TestCase):
    """End-to-end coverage through EligibleHouseholdSelectionService.generate,
    exercising _resolve_eligibility_rule, EligibleMember.is_eligible, and
    select_households together — including the no-benefitPlanCode path, which
    is the actual backward-compatibility guarantee for existing PWP
    deployments."""

    def test_no_benefit_plan_code_runs_the_quota_algorithm(self):
        groups = [
            _fake_group("g-female", "HH-F", "Female", 30, "Poorest"),
            _fake_group("g-youth", "HH-Y", "Male", 22, "Poorer"),
            _fake_group("g-other", "HH-O", "Male", 45, "Middle"),
            _fake_group(
                "g-not-fit",
                "HH-N",
                "Male",
                40,
                "Poorest",
                individual_json_ext={"fit_for_work": False},
            ),
        ]
        service = _FakeSelectionService(groups)

        selection_result, summary = service.generate(target_count=2)

        # The not-fit-for-work household has no eligible members at all under
        # the default (no benefitPlanCode) rule, so it's dropped entirely —
        # this is the actual backward-compatibility guarantee: unchanged
        # fit-for-work-only eligibility, running through the quota/reserve
        # algorithm exactly as before this change.
        self.assertEqual(summary["eligible_households"], 3)
        self.assertNotIn(
            "g-not-fit",
            {row.household.id for row in selection_result.selected},
        )
        self.assertEqual(summary["selected_households"], 2)
        self.assertEqual(summary["reserve_households"], 1)
        self.assertEqual(
            {row.household.id for row in selection_result.reserve},
            {"g-other"},
        )
        # Demographic categorization only happens under the quota algorithm.
        self.assertEqual(summary["selected_female_headed_households"], 1)
        self.assertEqual(summary["selected_youth_households"], 1)

    def test_benefit_plan_code_rmep_runs_the_simple_algorithm(self):
        groups = [
            _fake_group(
                "g-business",
                "HH-B",
                "Male",
                40,
                "Poorest",
                individual_json_ext={"data_source": "PWP", "business_experience": True},
            ),
            _fake_group(
                "g-no-business",
                "HH-NB",
                "Male",
                40,
                "Poorest",
                individual_json_ext={"data_source": "PWP", "business_experience": False},
            ),
            _fake_group(
                "g-wrong-source",
                "HH-WS",
                "Male",
                40,
                "Poorest",
                individual_json_ext={"data_source": "SCTP", "business_experience": True},
            ),
            _fake_group(
                "g-not-fit-but-pwp",
                "HH-NF",
                "Male",
                40,
                "Poorest",
                individual_json_ext={"data_source": "PWP", "fit_for_work": False},
            ),
        ]
        service = _FakeSelectionService(groups)

        selection_result, summary = service.generate(
            target_count=10,
            benefit_plan_code="RMEP",
        )

        selected_ids = [row.household.id for row in selection_result.main]
        # Wrong data source is excluded entirely; RMEP doesn't require
        # fit_for_work, so that PWP-sourced household still qualifies.
        self.assertEqual(
            set(selected_ids),
            {"g-business", "g-no-business", "g-not-fit-but-pwp"},
        )
        # Business-experienced households are ranked first.
        self.assertEqual(selected_ids[0], "g-business")
        # No reserve list and no demographic categorization under this strategy.
        self.assertEqual(selection_result.reserve, [])
        self.assertEqual(summary["reserve_households"], 0)
        self.assertEqual(summary["selected_female_headed_households"], 0)

    def test_benefit_plan_code_upg_requires_source_fitness_and_age_range(self):
        groups = [
            _fake_group(
                "g-eligible",
                "HH-E",
                "Male",
                30,
                "Poorest",
                individual_json_ext={"data_source": "SCTP", "fit_for_work": True},
            ),
            _fake_group(
                "g-too-old",
                "HH-TO",
                "Male",
                65,
                "Poorest",
                individual_json_ext={"data_source": "SCTP", "fit_for_work": True},
            ),
            _fake_group(
                "g-not-fit",
                "HH-NF",
                "Male",
                30,
                "Poorest",
                individual_json_ext={"data_source": "SCTP", "fit_for_work": False},
            ),
            _fake_group(
                "g-wrong-source",
                "HH-WS",
                "Male",
                30,
                "Poorest",
                individual_json_ext={"data_source": "PWP", "fit_for_work": True},
            ),
        ]
        service = _FakeSelectionService(groups)

        selection_result, summary = service.generate(
            target_count=10,
            benefit_plan_code="UPG",
        )

        self.assertEqual(
            {row.household.id for row in selection_result.main},
            {"g-eligible"},
        )
        self.assertEqual(summary["eligible_households"], 1)
        self.assertEqual(selection_result.reserve, [])


class HotspotAndMicroCatchmentResolutionTest(TestCase):
    @patch("household_validation.services.Hotspot.objects.filter")
    def test_resolve_hotspot_filters_by_uuid_code_and_validity(self, filter_mock):
        filter_mock.return_value.first.return_value = "hotspot-instance"
        service = EligibleHouseholdSelectionService()

        result = service._resolve_hotspot("hotspot-uuid", "HS01")

        self.assertEqual(result, "hotspot-instance")
        self.assertEqual(filter_mock.call_args[0][0], Q(uuid="hotspot-uuid") | Q(code="HS01"))
        self.assertEqual(filter_mock.call_args[1], {"validity_to__isnull": True})

    def test_resolve_hotspot_returns_none_without_identifiers(self):
        service = EligibleHouseholdSelectionService()
        self.assertIsNone(service._resolve_hotspot(None, None))

    @patch("household_validation.services.MicroCatchment.objects.filter")
    def test_resolve_micro_catchment_filters_by_uuid_code_and_validity(self, filter_mock):
        filter_mock.return_value.first.return_value = "catchment-instance"
        service = EligibleHouseholdSelectionService()

        result = service._resolve_micro_catchment("catchment-uuid", "MC01")

        self.assertEqual(result, "catchment-instance")
        self.assertEqual(filter_mock.call_args[0][0], Q(uuid="catchment-uuid") | Q(code="MC01"))
        self.assertEqual(filter_mock.call_args[1], {"validity_to__isnull": True})

    def test_resolve_micro_catchment_returns_none_without_identifiers(self):
        service = EligibleHouseholdSelectionService()
        self.assertIsNone(service._resolve_micro_catchment(None, None))


class LocationFilterHotspotAndCatchmentTest(TestCase):
    def test_scopes_to_hotspot_villages(self):
        service = EligibleHouseholdSelectionService()
        hotspot = SimpleNamespace(
            villages=SimpleNamespace(values_list=lambda *a, **k: ["V01", "V02"]),
        )
        queryset = MagicMock()

        with patch.object(service, "_resolve_hotspot", return_value=hotspot) as resolve_mock:
            result = service._apply_location_filters(queryset, hotspot_code="HS01")

        resolve_mock.assert_called_once_with(None, "HS01")
        queryset.filter.assert_called_once_with(location__code__in=["V01", "V02"])
        self.assertIs(result, queryset.filter.return_value)

    def test_unknown_hotspot_yields_no_households(self):
        service = EligibleHouseholdSelectionService()
        queryset = MagicMock()

        with patch.object(service, "_resolve_hotspot", return_value=None):
            result = service._apply_location_filters(queryset, hotspot_id="missing")

        queryset.none.assert_called_once()
        self.assertIs(result, queryset.none.return_value)

    def test_hotspot_with_no_villages_yields_no_households(self):
        service = EligibleHouseholdSelectionService()
        hotspot = SimpleNamespace(villages=SimpleNamespace(values_list=lambda *a, **k: []))
        queryset = MagicMock()

        with patch.object(service, "_resolve_hotspot", return_value=hotspot):
            result = service._apply_location_filters(queryset, hotspot_code="HS01")

        queryset.none.assert_called_once()
        self.assertIs(result, queryset.none.return_value)

    def test_explicit_village_selection_overrides_hotspot(self):
        service = EligibleHouseholdSelectionService()
        queryset = MagicMock()

        with patch.object(service, "_resolve_hotspot") as resolve_mock:
            service._apply_location_filters(
                queryset,
                village_codes=["V01"],
                hotspot_code="HS01",
            )

        resolve_mock.assert_not_called()

    def test_scopes_to_micro_catchment_hotspot_villages_before_broader_links(self):
        service = EligibleHouseholdSelectionService()
        micro_catchment = SimpleNamespace(
            traditional_authorities=MagicMock(),
            gvhs=MagicMock(),
        )
        queryset = MagicMock()
        hotspot_queryset = MagicMock()
        hotspot_queryset.values_list.return_value.distinct.return_value = ["V01", "V02"]

        with (
            patch.object(service, "_resolve_micro_catchment", return_value=micro_catchment) as resolve_mock,
            patch("household_validation.services.Hotspot.objects.filter", return_value=hotspot_queryset),
        ):
            result = service._apply_location_filters(queryset, catchment_code="MC01")

        resolve_mock.assert_called_once_with(None, "MC01")
        queryset.filter.assert_called_once_with(Q(location__code__in=["V01", "V02"]))
        micro_catchment.gvhs.filter.assert_not_called()
        micro_catchment.traditional_authorities.filter.assert_not_called()
        self.assertIs(result, queryset.filter.return_value)

    def test_micro_catchment_falls_back_to_active_gvhs_without_villages(self):
        service = EligibleHouseholdSelectionService()
        micro_catchment = SimpleNamespace(
            gvhs=MagicMock(),
            traditional_authorities=MagicMock(),
        )
        hotspot_queryset = MagicMock()
        hotspot_queryset.values_list.return_value.distinct.return_value = []
        micro_catchment.gvhs.filter.return_value.values_list.return_value = ["GVH01"]

        with patch(
            "household_validation.services.Hotspot.objects.filter",
            return_value=hotspot_queryset,
        ):
            location_filter = service._micro_catchment_location_filter(micro_catchment)

        self.assertEqual(
            location_filter,
            Q(location__code__in=["GVH01"])
            | Q(location__parent__code__in=["GVH01"]),
        )
        micro_catchment.traditional_authorities.filter.assert_not_called()

    def test_unknown_micro_catchment_yields_no_households(self):
        service = EligibleHouseholdSelectionService()
        queryset = MagicMock()

        with patch.object(service, "_resolve_micro_catchment", return_value=None):
            result = service._apply_location_filters(queryset, catchment_id="missing")

        queryset.none.assert_called_once()
        self.assertIs(result, queryset.none.return_value)

    def test_explicit_ta_selection_is_intersected_with_micro_catchment(self):
        service = EligibleHouseholdSelectionService()
        queryset = MagicMock()
        catchment_queryset = MagicMock()
        queryset.filter.return_value = catchment_queryset
        catchment_filter = Q(location__code__in=["V01"])

        with patch.object(
            service,
            "_micro_catchment_location_filter",
            return_value=catchment_filter,
        ):
            with patch.object(service, "_resolve_micro_catchment", return_value=SimpleNamespace()):
                result = service._apply_location_filters(
                    queryset,
                    ta_codes=["TA01"],
                    catchment_code="MC01",
                )

        queryset.filter.assert_called_once_with(catchment_filter)
        catchment_queryset.filter.assert_called_once_with(
            Q(location__code__in=["TA01"])
            | Q(location__parent__code__in=["TA01"])
            | Q(location__parent__parent__code__in=["TA01"])
        )
        self.assertIs(result, catchment_queryset.filter.return_value)


class CatchmentAndHotspotSelectionIntegrationTest(DjangoTestCase):
    """Prove GraphQL generation and preview exclude out-of-scope households."""

    @classmethod
    def setUpTestData(cls):
        cls.audit_user = User.objects.create(username="household-validation-scope-test")
        cls.district = Location.objects.create(
            code="HV-TEST-DISTRICT",
            name="Validation test district",
            type="R",
            audit_user_id=1,
        )
        cls.ta = Location.objects.create(
            code="HV-TEST-TA",
            name="Validation test TA",
            type="D",
            parent=cls.district,
            audit_user_id=1,
        )
        cls.gvh = Location.objects.create(
            code="HV-TEST-GVH",
            name="Validation test GVH",
            type="W",
            parent=cls.ta,
            audit_user_id=1,
        )
        cls.inside_village = Location.objects.create(
            code="HV-TEST-V-IN",
            name="Inside village",
            type="V",
            parent=cls.gvh,
            audit_user_id=1,
        )
        cls.outside_village = Location.objects.create(
            code="HV-TEST-V-OUT",
            name="Outside village",
            type="V",
            parent=cls.gvh,
            audit_user_id=1,
        )
        cls.catchment = MicroCatchment.objects.create(
            code="HV-TEST-MC",
            name="Validation test micro-catchment",
            district=cls.district,
            audit_user_id=1,
        )
        cls.hotspot = Hotspot.objects.create(
            code="HV-TEST-HS",
            name="Validation test hotspot",
            micro_catchment=cls.catchment,
            audit_user_id=1,
        )
        HotspotVillage.objects.create(
            hotspot=cls.hotspot,
            location=cls.inside_village,
            audit_user_id=1,
        )
        cls.inside_group = cls._create_eligible_household(
            "HV-HH-IN",
            cls.inside_village,
        )
        cls.outside_group = cls._create_eligible_household(
            "HV-HH-OUT",
            cls.outside_village,
        )

    @classmethod
    def _create_eligible_household(cls, code, village):
        group = Group(
            code=code,
            location=village,
            json_ext={
                "form_number": f"FORM-{code}",
                "household_wealth_quintile": "Poorest",
            },
        )
        group.save(user=cls.audit_user)
        individual = Individual(
            first_name=code,
            last_name="Worker",
            dob=date(1990, 1, 1),
            location=village,
            json_ext={"fit_for_work": True, "gender": "Male"},
        )
        individual.save(user=cls.audit_user)
        GroupIndividual.objects.bulk_create(
            [
                GroupIndividual(
                    id=uuid4(),
                    group=group,
                    individual=individual,
                    role=GroupIndividual.Role.HEAD,
                    user_created=cls.audit_user,
                    user_updated=cls.audit_user,
                )
            ]
        )
        return group

    def test_micro_catchment_excludes_district_households_from_generation_and_preview_without_ta(self):
        self._assert_generation_and_preview_scope(
            catchment_code=self.catchment.code,
        )

    def test_hotspot_excludes_households_in_other_villages_in_same_district_from_generation_and_preview(self):
        self._assert_generation_and_preview_scope(
            hotspot_code=self.hotspot.code,
        )

    def _assert_generation_and_preview_scope(self, **scope):
        info = SimpleNamespace(context=SimpleNamespace(user=self.audit_user))
        filters = {
            "district_id": self.district.id,
            "target_count": 10,
            **scope,
        }
        with (
            patch.object(
                Group,
                "get_queryset",
                side_effect=lambda queryset, user: queryset,
            ),
            patch.object(
                GenerateHouseholdValidationListMutation,
                "_validate_user",
            ),
            patch.object(Query, "_check_permissions"),
            patch(
                "household_validation.gql_mutations.HouseholdValidationProjectLookupService.list_projects",
                return_value=[],
            ),
            patch.object(
                EligibleHouseholdSelectionService,
                "_project_names",
                return_value=[],
            ),
        ):
            generated = GenerateHouseholdValidationListMutation.mutate(
                None,
                info,
                **filters,
            )
            preview = Query.resolve_household_validation_preview(
                None,
                info,
                **filters,
            )

        self.assertEqual(generated.total_households, 1)
        self.assertEqual(generated.selected_households, 1)
        self.assertEqual(len(generated.village_breakdown), 1)
        self.assertEqual(
            generated.village_breakdown[0]["village_code"],
            self.inside_village.code,
        )
        self.assertEqual(
            generated.village_breakdown[0]["selected_households"],
            1,
        )
        workbook = load_workbook(
            BytesIO(base64.b64decode(generated.file_base64)),
            read_only=True,
        )
        worksheet = workbook[VALIDATION_LIST_SHEET]
        headers = [cell.value for cell in worksheet[1]]
        group_uuid_column = headers.index("group_uuid") + 1
        generated_group_ids = {
            str(row[group_uuid_column - 1].value)
            for row in worksheet.iter_rows(min_row=2)
        }
        self.assertEqual(generated_group_ids, {str(self.inside_group.id)})
        self.assertNotIn(str(self.outside_group.id), generated_group_ids)

        preview_group_codes = {edge.node.group_code for edge in preview.edges}
        self.assertEqual(preview_group_codes, {self.inside_group.code})
        self.assertNotIn(self.outside_group.code, preview_group_codes)


class _FakeSelectionService(EligibleHouseholdSelectionService):
    def __init__(self, groups):
        super().__init__(user=None)
        self.groups = groups

    def _base_queryset(self):
        return _FakeQuerySet(self.groups)

    def _project_names(self, location):
        return []


class _FakeQuerySet:
    def __init__(self, groups):
        self.groups = groups

    def __iter__(self):
        return iter(self.groups)

    def filter(self, *args, **kwargs):
        return self

    def select_related(self, *args, **kwargs):
        return self

    def prefetch_related(self, *args, **kwargs):
        return self


class _FakeRelatedManager:
    def __init__(self, rows):
        self.rows = rows

    def all(self):
        return self.rows


def _fake_group(
    group_id,
    code,
    head_gender,
    worker_age,
    wealth_quintile,
    validation_status=None,
    individual_json_ext=None,
    recipient_type="PRIMARY",
):
    location = _fake_location_tree()
    individual_json = {
        "gender": head_gender,
        "fit_for_work": True,
        "household_wealth_quintile": wealth_quintile,
    }
    if individual_json_ext:
        individual_json.update(individual_json_ext)
    individual = SimpleNamespace(
        id=f"{group_id}-individual",
        first_name="Head",
        last_name="Person",
        dob=_dob_for_age(worker_age),
        json_ext=individual_json,
    )
    group_individual = SimpleNamespace(
        id=f"{group_id}-member",
        is_deleted=False,
        role="HEAD",
        recipient_type=recipient_type,
        individual=individual,
        individual_id=individual.id,
    )
    json_ext = {
        "head_id": individual.id,
        "household_wealth_quintile": wealth_quintile,
    }
    if validation_status:
        json_ext["validation_status"] = validation_status
    return SimpleNamespace(
        id=group_id,
        code=code,
        json_ext=json_ext,
        location=location,
        groupindividuals=_FakeRelatedManager([group_individual]),
    )


def _fake_location_tree():
    region = SimpleNamespace(type="R", name="Region", code="R01", parent=None, id=1)
    district = SimpleNamespace(type="D", name="District", code="D01", parent=region, id=2)
    ta = SimpleNamespace(type="W", name="Traditional Authority", code="TA01", parent=district, id=3)
    return SimpleNamespace(type="V", name="Village", code="V01", parent=ta, id=4)


class ProjectLookupTest(TestCase):
    def test_active_project_statuses_match_mvp_dropdown_scope(self):
        self.assertEqual(ACTIVE_PROJECT_STATUSES, ("PREPARATION", "IN_PROGRESS"))

    def test_project_option_from_project_serializes_project_fields(self):
        project = SimpleNamespace(
            id="project-1",
            name="Road Works",
            status="PREPARATION",
            location_id=12,
        )

        self.assertEqual(
            project_option_from_project(project),
            ProjectOption(
                id="project-1",
                name="Road Works",
                status="PREPARATION",
                location_id="12",
            ),
        )

    def test_project_option_from_project_falls_back_to_location_object(self):
        project = SimpleNamespace(
            id="project-2",
            name="Drainage",
            status="IN_PROGRESS",
            location=SimpleNamespace(id=21),
        )

        self.assertEqual(project_option_from_project(project).location_id, "21")


class GraphQLPermissionTest(TestCase):
    def test_require_permissions_allows_user_with_required_rights(self):
        user = SimpleNamespace(
            id="user-1",
            is_anonymous=False,
            has_perms=lambda perms: perms == ["958001"],
        )

        self.assertIsNone(require_permissions(user, ["958001"]))

    def test_require_permissions_rejects_missing_rights(self):
        user = SimpleNamespace(
            id="user-1",
            is_anonymous=False,
            has_perms=lambda perms: False,
        )

        with self.assertRaises(HouseholdValidationPermissionError):
            require_permissions(user, ["958002"])

    def test_require_permissions_rejects_anonymous_or_missing_user(self):
        anonymous_user = SimpleNamespace(
            id=None,
            is_anonymous=True,
            has_perms=lambda perms: True,
        )

        with self.assertRaises(HouseholdValidationPermissionError):
            require_permissions(anonymous_user, ["958001"])
        with self.assertRaises(HouseholdValidationPermissionError):
            require_permissions(None, ["958001"])


class ValidationUploadParserTest(TestCase):
    def test_business_header_aliases_and_whitespace(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        for column, label, value in (
            (HAS_BUSINESS_COLUMN, " business_experience ", "YES"),
            (BUSINESS_TYPE_COLUMN, "TYPE OF\nBUSINESS", "crop farming"),
            (BUSINESS_DURATION_COLUMN, "business_period", 2.5),
        ):
            index = EXCEL_COLUMNS.index(column) + 1
            worksheet.cell(1, index, label)
            worksheet.cell(2, index, value)
        parsed = parse_validation_workbook(self._workbook_bytes(workbook))
        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.rows[0].business_updates, {
            "business_experience": "Yes",
            "type_of_business": "Crop farming",
            "business_period": 2.5,
        })

    def test_duplicate_editable_header_is_rejected(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        worksheet.cell(1, worksheet.max_column + 1, "NATIONAL ID")
        parsed = parse_validation_workbook(self._workbook_bytes(workbook))
        self.assertEqual(parsed.errors, ["Duplicate column: national_id"])

    def test_invalid_business_values_block_household(self):
        for column, value in (
            (HAS_BUSINESS_COLUMN, "MAYBE"),
            (BUSINESS_TYPE_COLUMN, "Unknown business"),
            (BUSINESS_DURATION_COLUMN, -1),
            (BUSINESS_DURATION_COLUMN, 101),
            (BUSINESS_DURATION_COLUMN, "NaN"),
            (BUSINESS_DURATION_COLUMN, "forever"),
        ):
            with self.subTest(column=column, value=value):
                workbook = self._upload_workbook()
                worksheet = workbook[VALIDATION_LIST_SHEET]
                worksheet.cell(2, EXCEL_COLUMNS.index(HAS_BUSINESS_COLUMN) + 1, "Yes")
                worksheet.cell(2, EXCEL_COLUMNS.index(column) + 1, value)
                parsed = parse_validation_workbook(self._workbook_bytes(workbook))
                self.assertTrue(parsed.errors)
                self.assertEqual(parsed.rows, [])
                self.assertEqual(parsed.invalid_group_keys, frozenset({"group-1"}))

    def test_business_period_is_required_only_for_business_yes(self):
        for business, period, valid in (
            ("Yes", None, False), ("Yes", "  ", False),
            ("Yes", 0, True), ("Yes", 0.5, True),
            ("No", None, True), (None, None, True),
        ):
            with self.subTest(business=business, period=period):
                workbook = self._upload_workbook()
                sheet = workbook[VALIDATION_LIST_SHEET]
                sheet.cell(2, EXCEL_COLUMNS.index(HAS_BUSINESS_COLUMN) + 1, business)
                if business == "Yes":
                    sheet.cell(
                        2, EXCEL_COLUMNS.index(BUSINESS_TYPE_COLUMN) + 1,
                        "Crop farming",
                    )
                sheet.cell(2, EXCEL_COLUMNS.index(BUSINESS_DURATION_COLUMN) + 1, period)
                parsed = parse_validation_workbook(self._workbook_bytes(workbook))
                if valid:
                    self.assertEqual(parsed.errors, [])
                else:
                    self.assertEqual(parsed.rows, [])
                    self.assertIn("is required when", parsed.errors[0])
                    self.assertEqual(parsed.invalid_group_keys, frozenset({"group-1"}))

    def test_business_no_clears_old_details_and_blank_preserves_them(self):
        workbook = self._upload_workbook()
        parsed = parse_validation_workbook(self._workbook_bytes(workbook))
        self.assertEqual(parsed.rows[0].business_updates, {})
        worksheet = workbook[VALIDATION_LIST_SHEET]
        worksheet.cell(2, EXCEL_COLUMNS.index(HAS_BUSINESS_COLUMN) + 1, "No")
        parsed = parse_validation_workbook(self._workbook_bytes(workbook))
        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.rows[0].business_updates, {
            "business_experience": "No",
            "type_of_business": None,
            "business_period": None,
        })

    def test_parse_validation_workbook_resolves_project_name_to_hidden_project_id(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        project_column = EXCEL_COLUMNS.index("project") + 1
        primary_worker_column = EXCEL_COLUMNS.index("primary_worker") + 1
        worksheet.cell(row=2, column=project_column, value="Road Works")
        worksheet.cell(row=2, column=primary_worker_column, value="YES")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.rows_read, 1)
        self.assertEqual(parsed.rows[0].project_id, "project-1")
        self.assertIsNone(parsed.rows[0].verified)
        self.assertEqual(parsed.rows[0].primary_worker, True)
        self.assertEqual(PROJECT_SELECTION_TYPE_INTENT, "INTENT")

    def test_parse_validation_workbook_reports_missing_required_columns(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = VALIDATION_LIST_SHEET
        worksheet.append(["batch_id", "group_uuid"])

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.rows, [])
        self.assertTrue(parsed.errors[0].startswith("Missing required columns:"))
        self.assertIn("member_uuid", parsed.errors[0])

    def test_parse_validation_workbook_rejects_previous_schema(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        for column_number in sorted(
            (
                EXCEL_COLUMNS.index(column) + 1
                for column in OPTIONAL_UPLOAD_COLUMNS
            ),
            reverse=True,
        ):
            worksheet.delete_cols(column_number)
        worksheet.cell(
            row=1,
            column=worksheet.max_column + 1,
            value="current_recipient_type",
        )

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertTrue(parsed.errors[0].startswith("Missing required columns:"))

    def test_parse_validation_workbook_preserves_primary_worker_no_and_blank(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        primary_worker_column = EXCEL_COLUMNS.index("primary_worker") + 1
        worksheet.cell(row=2, column=primary_worker_column, value="NO")
        worksheet.append(
            [
                "batch-1" if column == "batch_id"
                else "group-2" if column == "group_uuid"
                else "member-2" if column == "member_uuid"
                else None
                for column in EXCEL_COLUMNS
            ]
        )

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.rows[0].primary_worker, False)
        self.assertIsNone(parsed.rows[1].primary_worker)

    def test_parse_validation_workbook_rejects_invalid_editable_values(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        primary_worker_column = EXCEL_COLUMNS.index("primary_worker") + 1
        worksheet.cell(row=2, column=primary_worker_column, value="LATER")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.rows, [])
        self.assertIn("Row 2: primary_worker must be YES or NO", parsed.errors)
        self.assertEqual(parsed.error_row_numbers, frozenset({2}))
        self.assertEqual(parsed.invalid_group_keys, frozenset({"group-1"}))
        self.assertEqual(parsed.rows_read, 1)

    def test_parse_validation_workbook_ignores_legacy_verified_column(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        legacy_column = worksheet.max_column + 1
        worksheet.cell(row=1, column=legacy_column, value="verified")
        worksheet.cell(row=2, column=legacy_column, value="MAYBE")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.rows_read, 1)
        self.assertIsNone(parsed.rows[0].verified)

    def test_parse_validation_workbook_rejects_unknown_project_names(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        project_column = EXCEL_COLUMNS.index("project") + 1
        worksheet.cell(row=2, column=project_column, value="Unknown Project")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.rows, [])
        self.assertIn("Row 2: project is not in the project options", parsed.errors)

    def test_parse_validation_workbook_rejects_hidden_project_id_without_project(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        project_id_column = EXCEL_COLUMNS.index("project_id") + 1
        worksheet.cell(row=2, column=project_id_column, value="project-1")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.rows, [])
        self.assertIn("Row 2: project_id cannot be set without project", parsed.errors)

    def test_parse_validation_workbook_rejects_hidden_project_id_mismatch(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        project_column = EXCEL_COLUMNS.index("project") + 1
        project_id_column = EXCEL_COLUMNS.index("project_id") + 1
        worksheet.cell(row=2, column=project_column, value="Road Works")
        worksheet.cell(row=2, column=project_id_column, value="different-project")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.rows, [])
        self.assertIn("Row 2: project is not in the project options", parsed.errors)

    def test_parse_validation_workbook_resolves_duplicate_project_name_labels(self):
        workbook = self._upload_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        project_options = workbook[PROJECT_OPTIONS_SHEET]
        project_options.cell(row=1, column=3, value="project_label")
        project_options.cell(row=2, column=3, value="Road Works (project-1)")
        project_options.append(["project-2", "Road Works", "Road Works (project-2)"])
        project_column = EXCEL_COLUMNS.index("project") + 1
        worksheet.cell(row=2, column=project_column, value="Road Works (project-2)")

        parsed = parse_validation_workbook(self._workbook_bytes(workbook))

        self.assertEqual(parsed.errors, [])
        self.assertEqual(parsed.rows[0].project_id, "project-2")
        self.assertEqual(parsed.rows[0].project_name, "Road Works")

    def _upload_workbook(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = VALIDATION_LIST_SHEET
        worksheet.append(EXCEL_COLUMNS)
        values = {column: None for column in EXCEL_COLUMNS}
        values.update(
            {
                "batch_id": "batch-1",
                "group_uuid": "group-1",
                "member_uuid": "member-1",
                "District": "District",
                "TA": "Traditional Authority",
                "GVH": "Group Village Head",
                "Village": "Village",
            }
        )
        worksheet.append([values[column] for column in EXCEL_COLUMNS])
        project_options = workbook.create_sheet(PROJECT_OPTIONS_SHEET)
        project_options.append(["project_id", "project"])
        project_options.append(["project-1", "Road Works"])
        return workbook

    def _workbook_bytes(self, workbook):
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()


class ExcelValidationListExporterTest(TestCase):
    def setUp(self):
        # Existing business collection tests exercise the Jobs-Now configuration.
        patcher = patch.object(HouseholdValidationConfig, "business_columns_enabled", True)
        patcher.start()
        self.addCleanup(patcher.stop)

    @patch.object(HouseholdValidationConfig, "business_columns_enabled", False)
    def test_disabled_business_columns_preserve_core_workbook_and_upload(self):
        from household_validation.excel import (
            PRIMARY_WORKER_ANSWERS_RANGE, PRIMARY_WORKER_NO_RANGE,
        )

        projects = [SimpleNamespace(id="project-1", name="Road Works")]
        content = ExcelValidationListExporter(
            self._selection_result(member_count=2), "batch-1", projects=projects,
        ).export_bytes()
        workbook = load_workbook(BytesIO(content))
        worksheet = workbook[VALIDATION_LIST_SHEET]
        headers = [cell.value for cell in worksheet[1]]
        self.assertEqual(headers, [col for col in EXCEL_COLUMNS if col not in BUSINESS_COLUMNS])
        self.assertNotIn(BUSINESS_TYPE_OPTIONS_SHEET, workbook.sheetnames)
        self.assertEqual(worksheet.max_row, 3)
        self.assertTrue(worksheet.protection.sheet)
        self.assertFalse(self._cell(worksheet, "validation_notes").protection.locked)
        self.assertEqual(self._value(worksheet, "member_uuid"), "member-1")
        self.assertEqual(self._value(worksheet, "national_id"), "NAT-001")
        self.assertEqual(self._cell(worksheet, "national_id").number_format, "@")
        for column in ("batch_id", "group_uuid", "member_uuid", "row_type", "project_id"):
            self.assertTrue(worksheet.column_dimensions[self._cell(worksheet, column).column_letter].hidden)
        for column in ("participant_status", "household_status"):
            cell = self._cell(worksheet, column)
            self.assertEqual(cell.data_type, "f")
            self.assertTrue(cell.protection.locked)
            self.assertNotIn("#REF!", cell.value)
        self.assertEqual(
            self._value(worksheet, "participant_status"),
            '=IF(OR(UPPER(TRIM(P2&""))="YES",UPPER(TRIM(P2&""))="Y",'
            'UPPER(TRIM(P2&""))="TRUE",UPPER(TRIM(P2&""))="1"),"VERIFIED","NOT_VERIFIED")',
        )
        participant_col = self._cell(worksheet, "participant_status").column_letter
        self.assertIn(f"${participant_col}$2:${participant_col}$3", self._value(worksheet, "household_status"))

        # Only the primary-worker and project dropdowns remain; no business rules.
        rules = worksheet.data_validations.dataValidation
        self.assertEqual(len(rules), 2)
        self.assertEqual(len(worksheet.conditional_formatting), 1)
        self.assertTrue(any("HouseholdPrimaryWorkerAnswers" in rule.formula1 for rule in rules))
        self.assertTrue(any(rule.formula1 == "'Project Options'!$C$2:$C$2" for rule in rules))
        self.assertEqual(set(workbook.defined_names), {PRIMARY_WORKER_ANSWERS_RANGE, PRIMARY_WORKER_NO_RANGE})
        for name, reference in (
            (PRIMARY_WORKER_ANSWERS_RANGE, "$E$2:$E$3"),
            (PRIMARY_WORKER_NO_RANGE, "$E$3"),
        ):
            self.assertEqual(list(workbook.defined_names[name].destinations), [(PROJECT_OPTIONS_SHEET, reference)])
        options = workbook[PROJECT_OPTIONS_SHEET]
        self.assertEqual(options.sheet_state, "hidden")
        self.assertEqual([options["E2"].value, options["E3"].value], ["YES", "NO"])

        self._cell(worksheet, "primary_worker").value = "YES"
        self._cell(worksheet, "validation_notes").value = "Confirmed"
        output = BytesIO()
        workbook.save(output)
        parsed = parse_validation_workbook(output.getvalue())
        self.assertEqual(parsed.errors, [])
        self.assertTrue(parsed.rows[0].primary_worker)
        self.assertEqual(parsed.rows[0].values["validation_notes"], "Confirmed")
        for column in BUSINESS_COLUMNS:
            self.assertIsNone(parsed.rows[0].values[column])

    def test_export_configuration_is_read_for_each_workbook(self):
        for enabled in (False, True, False):
            with self.subTest(enabled=enabled), patch.object(
                HouseholdValidationConfig, "business_columns_enabled", enabled,
            ):
                workbook = ExcelValidationListExporter(
                    self._selection_result(), "batch-1",
                ).export_workbook()
                headers = [cell.value for cell in workbook[VALIDATION_LIST_SHEET][1]]
                self.assertEqual(BUSINESS_COLUMNS.issubset(headers), enabled)
                self.assertEqual(BUSINESS_TYPE_OPTIONS_SHEET in workbook.sheetnames, enabled)
        self.assertTrue(BUSINESS_COLUMNS.issubset(EXCEL_COLUMNS))

    @patch.object(HouseholdValidationConfig, "business_columns_enabled", False)
    def test_empty_export_without_business_columns(self):
        content = ExcelValidationListExporter(
            SelectionResult(main=[], reserve=[]), "batch-1",
        ).export_bytes()
        workbook = load_workbook(BytesIO(content))
        worksheet = workbook[VALIDATION_LIST_SHEET]
        self.assertEqual(worksheet.max_row, 1)
        self.assertEqual(worksheet.max_column, len(EXCEL_COLUMNS) - 3)
        self.assertNotIn(BUSINESS_TYPE_OPTIONS_SHEET, workbook.sheetnames)
        self.assertEqual(len(worksheet.data_validations.dataValidation), 1)

    def test_upload_accepts_existing_business_workbooks_in_both_deployments(self):
        workbook = ExcelValidationListExporter(self._selection_result(), "batch-1").export_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        for column, value in {
            "primary_worker": "YES",
            HAS_BUSINESS_COLUMN: "Yes",
            BUSINESS_TYPE_COLUMN: "Crop farming",
            BUSINESS_DURATION_COLUMN: 2,
        }.items():
            self._cell(worksheet, column).value = value
        output = BytesIO()
        workbook.save(output)
        for enabled in (False, True):
            with self.subTest(enabled=enabled), patch.object(
                HouseholdValidationConfig, "business_columns_enabled", enabled,
            ):
                parsed = parse_validation_workbook(output.getvalue())
                self.assertEqual(parsed.errors, [])
                self.assertEqual(parsed.rows[0].values[BUSINESS_TYPE_COLUMN], "Crop farming")
                self.assertEqual(parsed.rows[0].values[BUSINESS_DURATION_COLUMN], 2)

    @patch("household_validation.services.IndividualService")
    def test_export_edit_upload_persists_member_fields(self, service_mock):
        workbook = ExcelValidationListExporter(
            self._selection_result(), batch_id="batch-1"
        ).export_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        for column, value in {
            "national_id": "00123456",
            "primary_worker": "YES",
            HAS_BUSINESS_COLUMN: "Yes",
            BUSINESS_TYPE_COLUMN: "Crop farming",
            BUSINESS_DURATION_COLUMN: 0.5,
            "validation_notes": "Visited and confirmed",
        }.items():
            worksheet.cell(2, EXCEL_COLUMNS.index(column) + 1, value)
        output = BytesIO()
        workbook.save(output)
        parsed = parse_validation_workbook(output.getvalue())
        self.assertEqual(parsed.errors, [])
        row = parsed.rows[0]
        self.assertTrue(row.primary_worker)
        member = SimpleNamespace(individual=SimpleNamespace(
            id="individual-1", json_ext={"existing": "preserved"},
        ))
        service = HouseholdValidationUploadService()
        with (
            patch.object(service, "_resolve_row", return_value=([], MagicMock(), member, None)),
            patch.object(service, "_save_batch_row"),
        ):
            errors, national_id_updated = service._apply_row(
                row, MagicMock(), date.today(), datetime.now(),
                allow_participant_update=True,
            )
        self.assertEqual(errors, [])
        self.assertTrue(national_id_updated)
        self.assertEqual(service_mock.return_value.update.call_args.args[0]["json_ext"], {
            "existing": "preserved",
            "national_id": "00123456",
            "business_experience": "Yes",
            "type_of_business": "Crop farming",
            "business_period": 0.5,
            "validation_notes": "Visited and confirmed",
        })
        self.assertFalse(service._apply_member_details(member, row))
        service_mock.reset_mock()
        with (
            patch.object(service, "_resolve_row", return_value=([], MagicMock(), member, None)),
            patch.object(service, "_save_batch_row"),
        ):
            service._apply_row(
                row, MagicMock(), date.today(), datetime.now(),
                allow_participant_update=False,
            )
        service_mock.return_value.update.assert_not_called()

    def test_export_workbook_writes_headers_and_member_rows(self):
        result = self._selection_result()

        workbook = ExcelValidationListExporter(result, batch_id="batch-1").export_workbook()
        worksheet = workbook["Validation List"]

        self.assertEqual(
            [worksheet.cell(row=1, column=index).value for index in range(1, len(EXCEL_COLUMNS) + 1)],
            EXCEL_COLUMNS,
        )
        self.assertIn("primary_worker", EXCEL_COLUMNS)
        self.assertNotIn("verified", EXCEL_COLUMNS)
        self.assertNotIn("head", EXCEL_COLUMNS)
        self.assertNotIn("validation_date", EXCEL_COLUMNS)
        self.assertNotIn("current_recipient_type", EXCEL_COLUMNS)
        form_number_index = EXCEL_COLUMNS.index("form_number")
        self.assertEqual(
            EXCEL_COLUMNS[form_number_index:form_number_index + 6],
            [
                "form_number",
                "member_name",
                "relationship",
                "member_dob",
                "national_id",
                "primary_worker",
            ],
        )
        self.assertNotIn("participant", EXCEL_COLUMNS)
        self.assertNotIn("group_code", EXCEL_COLUMNS)
        self.assertEqual(worksheet["A2"].value, "batch-1")
        self.assertEqual(self._value(worksheet, "row_type"), ROW_TYPE_MAIN)
        self.assertEqual(self._value(worksheet, "District"), "District")
        self.assertEqual(self._value(worksheet, "TA"), "Traditional Authority")
        self.assertEqual(self._value(worksheet, "GVH"), "Group Village Head")
        self.assertEqual(self._value(worksheet, "Village"), "Village")
        self.assertEqual(self._value(worksheet, "form_number"), "FORM-001")
        self.assertEqual(self._value(worksheet, "group_uuid"), "group-1")
        self.assertEqual(self._value(worksheet, "member_uuid"), "member-1")
        self.assertEqual(self._value(worksheet, "member_name"), "Ada Worker")
        self.assertEqual(self._value(worksheet, "national_id"), "NAT-001")
        self.assertEqual(self._value(worksheet, "marital_status"), "Married")
        self.assertEqual(self._value(worksheet, "disability"), "Not disabled")
        self.assertEqual(self._value(worksheet, "fit_for_work"), "YES")
        self.assertEqual(self._value(worksheet, "relationship"), "HEAD")
        self.assertEqual(self._value(worksheet, "pmt_score"), -1.234)
        self.assertEqual(
            self._value(worksheet, "household_wealth_quintile"),
            "Poorest",
        )

        group = result.main[0].household.source
        self.assertEqual(group.code, "HH-001")

    def test_export_workbook_writes_exact_relationship_roles(self):
        workbook = ExcelValidationListExporter(
            self._selection_result(member_count=3),
            batch_id="batch-1",
        ).export_workbook()
        worksheet = workbook["Validation List"]
        relationship_column = EXCEL_COLUMNS.index("relationship") + 1

        self.assertEqual(worksheet.cell(2, relationship_column).value, "HEAD")
        self.assertEqual(worksheet.cell(3, relationship_column).value, "SPOUSE")
        self.assertEqual(worksheet.cell(4, relationship_column).value, "SON")

    def test_export_hides_internal_ids_and_preserves_upload_matching(self):
        workbook = ExcelValidationListExporter(
            self._selection_result(member_count=2), batch_id="batch-1",
        ).export_workbook()
        content = BytesIO()
        workbook.save(content)
        workbook = load_workbook(BytesIO(content.getvalue()))
        worksheet = workbook[VALIDATION_LIST_SHEET]
        for column in ("batch_id", "group_uuid", "member_uuid", "row_type", "project_id"):
            cell = self._cell(worksheet, column)
            self.assertTrue(worksheet.column_dimensions[cell.column_letter].hidden)
            self.assertTrue(cell.protection.locked)
        for column in ("District", "form_number", "member_name", "national_id"):
            cell = self._cell(worksheet, column)
            self.assertFalse(worksheet.column_dimensions[cell.column_letter].hidden)
        parsed = parse_validation_workbook(content.getvalue())
        self.assertEqual(parsed.errors, [])
        self.assertEqual(len(parsed.rows), 2)
        self.assertEqual(parsed.rows[0].values["row_type"], ROW_TYPE_MAIN)
        self.assertEqual(parsed.rows[0].values["batch_id"], "batch-1")
        self.assertEqual(parsed.rows[0].values["group_uuid"], "group-1")
        self.assertEqual(parsed.rows[0].values["member_uuid"], "member-1")
        self.assertEqual(parsed.rows[1].values["member_uuid"], "member-2")

    def test_export_workbook_has_hidden_project_id_and_dropdowns(self):
        result = self._selection_result()
        projects = [SimpleNamespace(id="project-1", name="Road Works, Phase 1")]

        workbook = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
            projects=projects,
        ).export_workbook()
        worksheet = workbook["Validation List"]
        project_options = workbook[PROJECT_OPTIONS_SHEET]
        project_id_letter = worksheet.cell(1, EXCEL_COLUMNS.index("project_id") + 1).column_letter

        self.assertTrue(worksheet.column_dimensions[project_id_letter].hidden)
        self.assertEqual(project_options.sheet_state, "hidden")
        self.assertEqual(project_options["A2"].value, "project-1")
        self.assertEqual(project_options["B2"].value, "Road Works, Phase 1")
        self.assertEqual(project_options["C2"].value, "Road Works, Phase 1")
        formulas = {validation.formula1 for validation in worksheet.data_validations.dataValidation}
        self.assertTrue(any("HouseholdPrimaryWorkerAnswers" in formula for formula in formulas))
        self.assertIn("'Project Options'!$C$2:$C$2", formulas)

    def test_primary_worker_dropdown_prevents_second_yes_per_household(self):
        from household_validation.excel import (
            PRIMARY_WORKER_ANSWERS_RANGE, PRIMARY_WORKER_NO_RANGE,
            BUSINESS_TYPE_OPTIONS_SHEET,
        )

        workbook = ExcelValidationListExporter(
            self._selection_result(member_count=3), batch_id="batch-1",
        ).export_workbook()
        content = BytesIO()
        workbook.save(content)
        workbook = load_workbook(BytesIO(content.getvalue()))
        worksheet = workbook[VALIDATION_LIST_SHEET]
        worker = self._cell(worksheet, "primary_worker").column_letter
        group = self._cell(worksheet, "group_uuid").column_letter
        rule = next(v for v in worksheet.data_validations.dataValidation
                    if f"{worker}2" in v.sqref)
        self.assertEqual(str(rule.sqref), f"{worker}2:{worker}4")
        self.assertEqual(rule.formula1,
            f'INDIRECT(IF(COUNTIFS(${group}$2:${group}$4,${group}2,'
            f'${worker}$2:${worker}$4,"YES")-COUNTIF(${worker}2,"YES")=0,'
            f'"{PRIMARY_WORKER_ANSWERS_RANGE}","{PRIMARY_WORKER_NO_RANGE}"))')
        self.assertLessEqual(len(rule.formula1), 255)
        self.assertTrue(rule.showErrorMessage)
        self.assertEqual(rule.errorStyle, "stop")
        self.assertTrue(rule.showInputMessage)
        self.assertTrue(rule.allow_blank)
        self.assertEqual(list(workbook.defined_names[PRIMARY_WORKER_NO_RANGE].destinations),
                         [(BUSINESS_TYPE_OPTIONS_SHEET, "$E$3")])
        options = workbook[BUSINESS_TYPE_OPTIONS_SHEET]
        self.assertEqual([options["E2"].value, options["E3"].value], ["YES", "NO"])
        self.assertTrue(any(f"{worker}2" in item.sqref
                            for item in worksheet.conditional_formatting))

    def test_business_dropdown_requires_primary_worker_answer(self):
        from household_validation.excel import (
            BUSINESS_ANSWERS_RANGE, BUSINESS_UNAVAILABLE_RANGE,
            BUSINESS_TYPE_OPTIONS_SHEET,
        )

        workbook = ExcelValidationListExporter(
            self._selection_result(member_count=2), batch_id="batch-1",
        ).export_workbook()
        content = BytesIO()
        workbook.save(content)
        workbook = load_workbook(BytesIO(content.getvalue()))
        worksheet = workbook[VALIDATION_LIST_SHEET]
        column = self._cell(worksheet, HAS_BUSINESS_COLUMN).column_letter
        worker = self._cell(worksheet, "primary_worker").column_letter
        validation = next(
            item for item in worksheet.data_validations.dataValidation
            if f"{column}2" in item.sqref
        )
        self.assertEqual(str(validation.sqref), f"{column}2:{column}3")
        self.assertIn(f'TRIM(${worker}2&"")="YES"', validation.formula1)
        self.assertNotIn('="NO"', validation.formula1)
        self.assertLessEqual(len(validation.formula1), 255)
        self.assertTrue(validation.showInputMessage)
        self.assertEqual(validation.promptTitle, "Select Primary Worker first")
        self.assertTrue(validation.showErrorMessage)
        self.assertEqual(validation.errorStyle, "stop")
        self.assertFalse(validation.allow_blank)
        self.assertFalse(validation.showDropDown)
        options = workbook[BUSINESS_TYPE_OPTIONS_SHEET]
        self.assertEqual(options.sheet_state, "hidden")
        self.assertEqual([options["C2"].value, options["C3"].value], ["Yes", "No"])
        self.assertEqual(options["C4"].value, '=""')
        self.assertEqual(list(workbook.defined_names[BUSINESS_ANSWERS_RANGE].destinations),
                         [(BUSINESS_TYPE_OPTIONS_SHEET, "$C$2:$C$3")])
        self.assertEqual(list(workbook.defined_names[BUSINESS_UNAVAILABLE_RANGE].destinations),
                         [(BUSINESS_TYPE_OPTIONS_SHEET, "$C$4")])

    def test_business_period_has_required_prompt_and_missing_value_highlight(self):
        workbook = ExcelValidationListExporter(
            self._selection_result(), batch_id="batch-1",
        ).export_workbook()
        content = BytesIO()
        workbook.save(content)
        worksheet = load_workbook(BytesIO(content.getvalue()))[VALIDATION_LIST_SHEET]
        cell = self._cell(worksheet, BUSINESS_DURATION_COLUMN)
        validation = next(v for v in worksheet.data_validations.dataValidation
                          if cell.coordinate in v.sqref)
        self.assertFalse(validation.allow_blank)
        self.assertEqual(validation.errorStyle, "stop")
        self.assertTrue(validation.showInputMessage)
        self.assertTrue(validation.showErrorMessage)
        self.assertEqual(validation.promptTitle, "Select Business Type first")
        self.assertTrue(any(cell.coordinate in rule.sqref
                            for rule in worksheet.conditional_formatting))

    def test_export_workbook_disambiguates_duplicate_project_names(self):
        result = self._selection_result()
        projects = [
            SimpleNamespace(id="project-1", name="Road Works"),
            SimpleNamespace(id="project-2", name="Road Works"),
        ]

        workbook = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
            projects=projects,
        ).export_workbook()
        project_options = workbook[PROJECT_OPTIONS_SHEET]

        self.assertEqual(project_options["C2"].value, "Road Works (project-1)")
        self.assertEqual(project_options["C3"].value, "Road Works (project-2)")

    def test_export_workbook_locks_structural_cells_and_unlocks_field_inputs(self):
        workbook = ExcelValidationListExporter(
            self._selection_result(),
            batch_id="batch-1",
        ).export_workbook()
        worksheet = workbook["Validation List"]

        self.assertTrue(worksheet.protection.sheet)
        for column in (
            "group_uuid",
            "relationship",
            "household_wealth_quintile",
        ):
            self.assertTrue(self._cell(worksheet, column).protection.locked)
        for column in (
            "national_id",
            "primary_worker",
            "validation_notes",
        ):
            self.assertFalse(self._cell(worksheet, column).protection.locked)

    def test_export_workbook_writes_one_row_per_selected_eligible_member(self):
        result = self._selection_result(member_count=2)

        workbook = ExcelValidationListExporter(result, batch_id="batch-1").export_workbook()
        worksheet = workbook["Validation List"]

        self.assertEqual(worksheet.max_row, 3)
        self.assertEqual(self._value(worksheet, "member_uuid", 2), "member-1")
        self.assertEqual(self._value(worksheet, "member_uuid", 3), "member-2")

    def test_export_workbook_alternates_row_colours_by_household(self):
        first = self._selection_result(member_count=2)
        second = self._selection_result(member_count=1)
        third = self._selection_result(member_count=1)
        second.main[0].household.id = "group-2"
        third.main[0].household.id = "group-3"
        result = SelectionResult(
            main=[*first.main, *second.main, *third.main],
            reserve=[],
        )

        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        self.assertEqual(worksheet["A2"].fill.fgColor.rgb, HOUSEHOLD_ROW_COLORS[0])
        self.assertEqual(worksheet["A3"].fill.fgColor.rgb, HOUSEHOLD_ROW_COLORS[0])
        self.assertEqual(worksheet["A4"].fill.fgColor.rgb, HOUSEHOLD_ROW_COLORS[1])
        self.assertEqual(worksheet["A5"].fill.fgColor.rgb, HOUSEHOLD_ROW_COLORS[0])
        for row_number in range(2, 6):
            expected_fill = worksheet.cell(row_number, 1).fill.fgColor.rgb
            self.assertTrue(
                all(
                    cell.fill.fgColor.rgb == expected_fill
                    for cell in worksheet[row_number]
                )
            )

    def test_export_workbook_preserves_identifier_columns_as_text(self):
        result = self._selection_result()
        member = result.main[0].household.eligible_members[0]
        individual = member.source.individual
        individual.json_ext["form_number"] = "001234"
        individual.json_ext["national_id"] = "00123456"

        exported = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_bytes()
        worksheet = load_workbook(BytesIO(exported))["Validation List"]

        for column, expected in (
            ("form_number", "001234"),
            ("national_id", "00123456"),
        ):
            cell = self._cell(worksheet, column)
            self.assertEqual(cell.value, expected)
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(cell.number_format, "@")

    def test_export_starts_business_fields_and_notes_blank_despite_stored_answers(self):
        result = self._selection_result(member_count=2)
        for member in result.main[0].household.eligible_members:
            member.source.individual.json_ext.update({
                "business_experience": "Yes",
                "type_of_business": "Crop farming",
                "business_period": 3,
                "validation_notes": "Previous visit notes",
            })
        workbook = ExcelValidationListExporter(result, batch_id="batch-1").export_workbook()
        content = BytesIO()
        workbook.save(content)
        worksheet = load_workbook(BytesIO(content.getvalue()))[VALIDATION_LIST_SHEET]
        for row in range(2, worksheet.max_row + 1):
            for column in ("primary_worker", HAS_BUSINESS_COLUMN, BUSINESS_TYPE_COLUMN,
                           BUSINESS_DURATION_COLUMN, "validation_notes"):
                self.assertIsNone(self._value(worksheet, column, row))
                self.assertFalse(self._cell(worksheet, column, row).protection.locked)
        # Generating a fresh form does not erase the saved answers.
        for member in result.main[0].household.eligible_members:
            self.assertEqual(member.source.individual.json_ext["business_period"], 3)
            self.assertEqual(member.source.individual.json_ext["validation_notes"],
                             "Previous visit notes")

    def test_export_workbook_leaves_primary_worker_blank(self):
        result = self._selection_result(member_count=3)
        stored_values = (
            (True, "SECONDARY"),
            (False, "PRIMARY"),
            (None, "PRIMARY"),
        )
        for selected_member, (stored_value, recipient_type) in zip(
            result.main[0].household.eligible_members,
            stored_values,
        ):
            selected_member.source.json_ext = (
                {"primary_worker": stored_value}
                if stored_value is not None
                else {}
            )
            selected_member.source.recipient_type = recipient_type

        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        self.assertIsNone(self._value(worksheet, "primary_worker", 2))
        self.assertIsNone(self._value(worksheet, "primary_worker", 3))
        self.assertIsNone(self._value(worksheet, "primary_worker", 4))

    def test_export_workbook_resolves_micro_catchment_from_gvh_link(self):
        micro_catchment = SimpleNamespace(name="Catchment A", code="MC-A")
        link = SimpleNamespace(micro_catchment=micro_catchment)
        gvh_manager = MagicMock()
        gvh_manager.filter.return_value.select_related.return_value.first.return_value = link

        district = SimpleNamespace(type="R", name="District", code="D01", parent=None)
        ta = SimpleNamespace(
            type="D", name="Traditional Authority", code="TA01", parent=district, id=2,
        )
        gvh = SimpleNamespace(
            type="W", name="Group Village Head", code="GVH01", parent=ta, id=3,
            micro_catchments_gvh=gvh_manager,
        )
        village = SimpleNamespace(type="V", name="Village", code="V01", parent=gvh)

        result = self._selection_result(location=village)
        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        self.assertEqual(self._value(worksheet, MICRO_CATCHMENT_COLUMN), "Catchment A")
        gvh_manager.filter.assert_called_once_with(
            validity_to__isnull=True,
            micro_catchment__validity_to__isnull=True,
        )

    def test_export_workbook_falls_back_to_ta_micro_catchment_link(self):
        micro_catchment = SimpleNamespace(name=None, code="MC-B")
        link = SimpleNamespace(micro_catchment=micro_catchment)
        ta_manager = MagicMock()
        ta_manager.filter.return_value.select_related.return_value.first.return_value = link
        gvh_manager = MagicMock()
        gvh_manager.filter.return_value.select_related.return_value.first.return_value = None

        district = SimpleNamespace(type="R", name="District", code="D01", parent=None)
        ta = SimpleNamespace(
            type="D", name="Traditional Authority", code="TA01", parent=district, id=2,
            micro_catchments_ta=ta_manager,
        )
        gvh = SimpleNamespace(
            type="W", name="Group Village Head", code="GVH01", parent=ta, id=3,
            micro_catchments_gvh=gvh_manager,
        )
        village = SimpleNamespace(type="V", name="Village", code="V01", parent=gvh)

        result = self._selection_result(location=village)
        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        # No name on the micro-catchment, so the code is used as the label.
        self.assertEqual(self._value(worksheet, MICRO_CATCHMENT_COLUMN), "MC-B")

    def test_export_workbook_leaves_micro_catchment_blank_without_a_link(self):
        result = self._selection_result()

        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        self.assertIsNone(self._value(worksheet, MICRO_CATCHMENT_COLUMN))

    def test_export_workbook_resolves_hotspot_from_village_link(self):
        hotspot = SimpleNamespace(name="Hotspot A", code="HS-A")
        link = SimpleNamespace(hotspot=hotspot)
        village_manager = MagicMock()
        village_manager.filter.return_value.select_related.return_value.first.return_value = link

        district = SimpleNamespace(type="R", name="District", code="D01", parent=None)
        ta = SimpleNamespace(type="D", name="Traditional Authority", code="TA01", parent=district)
        gvh = SimpleNamespace(type="W", name="Group Village Head", code="GVH01", parent=ta)
        village = SimpleNamespace(
            type="V", name="Village", code="V01", parent=gvh, id=4,
            hotspot_links=village_manager,
        )

        result = self._selection_result(location=village)
        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        self.assertEqual(self._value(worksheet, HOTSPOT_COLUMN), "Hotspot A")
        village_manager.filter.assert_called_once_with(
            validity_to__isnull=True,
            hotspot__validity_to__isnull=True,
        )

    def test_export_workbook_leaves_hotspot_blank_without_a_link(self):
        result = self._selection_result()

        worksheet = ExcelValidationListExporter(
            result,
            batch_id="batch-1",
        ).export_workbook()["Validation List"]

        self.assertIsNone(self._value(worksheet, HOTSPOT_COLUMN))

    def _selection_result(self, member_count=1, location=None):
        location = location or self._location_tree()
        group = SimpleNamespace(id="group-1", code="HH-001", location=location)
        selected_members = []
        for index in range(1, member_count + 1):
            individual = SimpleNamespace(
                first_name="Ada" if index == 1 else "Grace",
                last_name="Worker",
                json_ext={
                    "form_number": "FORM-001",
                    "national_id": f"NAT-{index:03d}",
                    "marital_status": "Married" if index == 1 else "Never married",
                    "disability": "Not disabled",
                },
            )
            group_individual = SimpleNamespace(individual=individual)
            selected_members.append(
                EligibleMember(
                    id=f"member-{index}",
                    gender="Female",
                    dob=date(1990, 1, index),
                    fit_for_work=True,
                    role=("HEAD" if index == 1 else "SPOUSE" if index == 2 else "SON"),
                    recipient_type="PRIMARY" if index == 1 else None,
                    source=group_individual,
                )
            )
        household = EligibleHousehold(
            id="group-1",
            code="HH-001",
            wealth_quintile="Poorest",
            pmt_score=-1.234,
            head=selected_members[0],
            eligible_members=selected_members,
            source=group,
        )
        return SelectionResult(
            main=[
                SelectedHousehold(
                    household=household,
                    category=CATEGORY_FEMALE_HEADED,
                    row_type=ROW_TYPE_MAIN,
                )
            ],
            reserve=[],
        )

    def _cell(self, worksheet, column_name, row_number=2):
        headers = [cell.value for cell in worksheet[1]]
        return worksheet.cell(row_number, headers.index(column_name) + 1)

    def _value(self, worksheet, column_name, row_number=2):
        return self._cell(worksheet, column_name, row_number).value

    def _location_tree(self):
        district = SimpleNamespace(type="R", name="District", code="D01", parent=None)
        ta = SimpleNamespace(type="D", name="Traditional Authority", code="TA01", parent=district)
        gvh = SimpleNamespace(type="W", name="Group Village Head", code="GVH01", parent=ta)
        return SimpleNamespace(type="V", name="Village", code="V01", parent=gvh)


class UploadHardeningTest(TestCase):
    @staticmethod
    def _uploaded_row(row_number, member_uuid, primary_worker):
        return UploadedValidationRow(
            row_number=row_number,
            values={
                "group_uuid": "group-1",
                "member_uuid": member_uuid,
                "form_number": "FORM-001",
                HAS_BUSINESS_COLUMN: "No",
            },
            verified=True,
            primary_worker=primary_worker,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )

    @staticmethod
    def _group_with_primary_workers(*primary_worker_values):
        members = [
            SimpleNamespace(
                individual_id=f"member-{index}",
                is_deleted=False,
                json_ext={"primary_worker": primary_worker},
            )
            for index, primary_worker in enumerate(primary_worker_values, start=1)
        ]
        return SimpleNamespace(
            id="group-1",
            code="HH-001",
            _prefetched_objects_cache={"groupindividuals": members},
        )

    def test_primary_worker_preflight_ignores_stored_value_for_blank_cell(self):
        group = self._group_with_primary_workers(True, False)
        rows = [
            self._uploaded_row(2, "member-1", None),
            self._uploaded_row(3, "member-2", True),
        ]
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_structural_errors", return_value=[]),
        ):
            rejections = service._primary_worker_rejections(rows)

        self.assertEqual(rejections, {})

    def test_primary_worker_preflight_accepts_explicit_worker_transfer(self):
        group = self._group_with_primary_workers(True, False)
        rows = [
            self._uploaded_row(2, "member-1", False),
            self._uploaded_row(3, "member-2", True),
        ]
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_structural_errors", return_value=[]),
        ):
            rejections = service._primary_worker_rejections(rows)

        self.assertEqual(rejections, {})

    def test_primary_worker_status_is_verified_with_exactly_one_worker(self):
        service = HouseholdValidationUploadService()
        rows = [self._uploaded_row(2, "member-1", True)]

        with patch.object(
            service,
            "_projected_primary_workers",
            return_value={"member-1": True, "member-2": False},
        ):
            statuses = service._primary_worker_verification_statuses(
                rows,
                eligible_group_keys={"group-1"},
            )

        self.assertEqual(statuses, {"group-1": True})

    def test_primary_worker_status_is_not_verified_without_a_worker(self):
        service = HouseholdValidationUploadService()
        rows = [self._uploaded_row(2, "member-1", False)]

        with patch.object(
            service,
            "_projected_primary_workers",
            return_value={"member-1": False, "member-2": False},
        ):
            statuses = service._primary_worker_verification_statuses(
                rows,
                eligible_group_keys={"group-1"},
            )

        self.assertEqual(statuses, {"group-1": False})

    def test_primary_worker_status_skips_rejected_household(self):
        service = HouseholdValidationUploadService()
        rows = [self._uploaded_row(2, "member-1", True)]

        statuses = service._primary_worker_verification_statuses(
            rows,
            eligible_group_keys={"group-1"},
            rejected_group_keys={"group-1"},
        )

        self.assertEqual(statuses, {})

    def test_primary_worker_preflight_ignores_ineligible_households(self):
        group = self._group_with_primary_workers(True, False)
        rows = [
            UploadedValidationRow(
                row_number=2,
                values={
                    "group_uuid": "group-1",
                    "member_uuid": "member-2",
                    "form_number": "FORM-001",
                },
                verified=None,
                primary_worker=True,
                validation_date=None,
                project_name=None,
                project_id=None,
                notes=None,
            ),
        ]
        service = HouseholdValidationUploadService()

        with patch.object(service, "_group", return_value=group):
            rejections = service._primary_worker_rejections(
                rows,
                eligible_group_keys=set(),
            )

        self.assertEqual(rejections, {})

    @patch("household_validation.services.parse_validation_workbook")
    def test_dry_run_reports_unique_rejected_household_without_writes(
        self,
        parse_workbook_mock,
    ):
        rows = [
            self._uploaded_row(2, "member-1", True),
            self._uploaded_row(3, "member-2", True),
        ]
        parse_workbook_mock.return_value = SimpleNamespace(
            rows=rows,
            rows_read=2,
            errors=[],
        )
        group = self._group_with_primary_workers(False, False)
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_structural_errors", return_value=[]),
            patch.object(service, "_get_or_create_batch") as create_batch_mock,
        ):
            result = service.upload(b"workbook", dry_run=True)

        self.assertEqual(result["households_with_multiple_primary_workers"], 1)
        self.assertEqual(result["participants_rejected"], 2)
        self.assertEqual(result["participants_not_verified"], 0)
        self.assertEqual(result["errors"], 0)
        self.assertEqual(result["error_messages"], [])
        create_batch_mock.assert_not_called()

    @patch("household_validation.services.parse_validation_workbook")
    def test_dry_run_keeps_parse_errors_separate_from_rejections(
        self,
        parse_workbook_mock,
    ):
        rows = [
            self._uploaded_row(2, "member-1", True),
            self._uploaded_row(3, "member-2", True),
        ]
        parse_workbook_mock.return_value = SimpleNamespace(
            rows=rows,
            rows_read=2,
            errors=["Row 2: invalid value"],
            error_row_numbers=frozenset({2}),
        )
        group = self._group_with_primary_workers(False, False)
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_structural_errors", return_value=[]),
        ):
            result = service.upload(b"workbook", dry_run=True)

        self.assertEqual(result["participants_rejected"], 2)
        self.assertEqual(result["errors"], 1)

    def test_primary_worker_rejection_explains_reason(self):
        service = HouseholdValidationUploadService()
        uploaded_row = self._uploaded_row(2, "member-1", True)

        with (
            patch.object(service, "_group", return_value=None),
            patch.object(service, "_group_individual", return_value=None),
            patch.object(service, "_project", return_value=None),
            patch.object(service, "_save_batch_row") as save_row_mock,
        ):
            service._save_primary_worker_rejection(
                uploaded_row,
                batch=MagicMock(),
            )

        self.assertEqual(
            save_row_mock.call_args.kwargs["error_message"],
            "household has more than one primary worker",
        )
        self.assertEqual(
            save_row_mock.call_args.kwargs["error_code"],
            "MULTIPLE_PRIMARY_WORKERS",
        )
        self.assertEqual(
            save_row_mock.call_args.kwargs["status"],
            HouseholdValidationBatchRow.Status.REJECTED,
        )
        self.assertEqual(
            save_row_mock.call_args.kwargs["uploaded_row"].participant_status,
            "REJECTED",
        )

    @patch("household_validation.services.HouseholdValidationBatchRow")
    def test_saved_rows_are_scoped_to_current_upload_attempt(self, batch_row_mock):
        service = HouseholdValidationUploadService()
        service._upload_attempt_id = uuid4()

        service._save_batch_row(
            batch=MagicMock(),
            uploaded_row=self._uploaded_row(2, "member-1", True),
        )

        self.assertEqual(
            batch_row_mock.call_args.kwargs["upload_attempt_id"],
            service._upload_attempt_id,
        )

    @patch("household_validation.services.parse_validation_workbook")
    def test_upload_rejects_every_row_without_partial_household_updates(
        self,
        parse_workbook_mock,
    ):
        rows = [
            self._uploaded_row(2, "member-1", True),
            self._uploaded_row(3, "member-2", True),
        ]
        parse_workbook_mock.return_value = SimpleNamespace(
            rows=rows,
            rows_read=2,
            errors=[],
        )
        group = self._group_with_primary_workers(False, False)
        batch = MagicMock()
        batch.id = "batch-1"
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_structural_errors", return_value=[]),
            patch.object(service, "_get_or_create_batch", return_value=batch),
            patch.object(
                service,
                "_save_primary_worker_rejection",
            ) as reject_mock,
            patch.object(service, "_apply_row") as apply_row_mock,
            patch(
                "household_validation.services.transaction.atomic",
                return_value=nullcontext(),
            ),
        ):
            result = service.upload(b"workbook")

        self.assertEqual(result["households_with_multiple_primary_workers"], 1)
        self.assertEqual(result["participants_rejected"], 2)
        self.assertEqual(result["errors"], 0)
        self.assertEqual(result["households_verified"], 0)
        self.assertEqual(result["households_not_verified"], 0)
        self.assertEqual(result["batch_id"], "batch-1")
        self.assertEqual(result["upload_attempt_id"], service._upload_attempt_id)
        self.assertIsNotNone(result["upload_attempt_id"])
        self.assertEqual(reject_mock.call_count, 2)
        apply_row_mock.assert_not_called()

    @patch("household_validation.services.GroupIndividualService")
    def test_primary_worker_update_preserves_json_and_recipient_type(
        self,
        service_class_mock,
    ):
        group_individual = SimpleNamespace(
            id="membership-1",
            group_id="group-1",
            recipient_type="SECONDARY",
            json_ext={"existing": "value"},
        )

        service = HouseholdValidationUploadService()
        service._apply_primary_worker(group_individual, True)
        service._apply_primary_worker(group_individual, False)

        service_class_mock.return_value.update.assert_has_calls(
            [
                call(
                    {
                        "id": "membership-1",
                        "group_id": "group-1",
                        "json_ext": {
                            "existing": "value",
                            "primary_worker": True,
                        },
                    }
                ),
                call(
                    {
                        "id": "membership-1",
                        "group_id": "group-1",
                        "json_ext": {
                            "existing": "value",
                            "primary_worker": False,
                        },
                    }
                ),
            ]
        )
        self.assertEqual(group_individual.recipient_type, "SECONDARY")

    @patch("household_validation.services.IndividualService")
    def test_national_id_update_preserves_individual_json(self, service_class_mock):
        individual = SimpleNamespace(
            id="individual-1",
            json_ext={"existing": "value", "national_id": "OLD-001"},
        )
        group_individual = SimpleNamespace(individual=individual)
        service = HouseholdValidationUploadService()

        service._apply_national_id(group_individual, " NEW-002 ")

        service_class_mock.return_value.update.assert_called_once_with(
            {
                "id": "individual-1",
                "json_ext": {
                    "existing": "value",
                    "national_id": "NEW-002",
                },
            }
        )

    def test_apply_row_counts_changed_national_id_as_participant_update(self):
        group = SimpleNamespace(id="group-1")
        group_individual = SimpleNamespace(
            individual=SimpleNamespace(
                id="individual-1",
                json_ext={"national_id": "OLD-001"},
            ),
        )
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={
                "group_uuid": "group-1",
                "member_uuid": "member-1",
                "form_number": "FORM-001",
                "national_id": "NEW-002",
            },
            verified=None,
            primary_worker=None,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )
        service = HouseholdValidationUploadService()

        with (
            patch.object(
                service,
                "_resolve_row",
                return_value=([], group, group_individual, None),
            ),
            patch.object(service, "_apply_national_id") as apply_national_id_mock,
            patch.object(service, "_save_batch_row") as save_row_mock,
        ):
            errors, participant_updated = service._apply_row(
                uploaded_row,
                batch=MagicMock(),
                upload_date=date.today(),
                uploaded_at=datetime.now(),
                allow_participant_update=True,
            )

        self.assertEqual(errors, [])
        self.assertTrue(participant_updated)
        apply_national_id_mock.assert_called_once_with(
            group_individual,
            "NEW-002",
        )
        self.assertEqual(
            save_row_mock.call_args.kwargs["status"],
            HouseholdValidationBatchRow.Status.APPLIED,
        )

    def test_apply_row_preserves_primary_worker_for_blank_cell(self):
        group = SimpleNamespace(id="group-1")
        group_individual = SimpleNamespace(
            individual_id="member-1",
            json_ext={"primary_worker": True},
        )
        service = HouseholdValidationUploadService()

        def uploaded_row(primary_worker):
            return UploadedValidationRow(
                row_number=2,
                values={
                    "group_uuid": "group-1",
                    "member_uuid": "member-1",
                    "form_number": "FORM-001",
                },
                verified=None,
                primary_worker=primary_worker,
                validation_date=None,
                project_name=None,
                project_id=None,
                notes=None,
            )

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(
                service,
                "_group_individual",
                return_value=group_individual,
            ),
            patch.object(service, "_project", return_value=None),
            patch.object(service, "_structural_errors", return_value=[]),
            patch.object(service, "_apply_primary_worker") as apply_worker_mock,
            patch.object(service, "_save_batch_row") as save_row_mock,
        ):
            unchanged_errors, unchanged = service._apply_row(
                uploaded_row(True),
                batch=MagicMock(),
                upload_date=date.today(),
                uploaded_at=datetime.now(),
                allow_participant_update=True,
            )
            blank_errors, blank = service._apply_row(
                uploaded_row(None),
                batch=MagicMock(),
                upload_date=date.today(),
                uploaded_at=datetime.now(),
                allow_participant_update=True,
            )

        self.assertEqual(unchanged_errors, [])
        self.assertFalse(unchanged)
        self.assertEqual(blank_errors, [])
        self.assertFalse(blank)
        apply_worker_mock.assert_not_called()
        self.assertEqual(
            [
                recorded_call.kwargs["status"]
                for recorded_call in save_row_mock.call_args_list
            ],
            [
                HouseholdValidationBatchRow.Status.SKIPPED,
                HouseholdValidationBatchRow.Status.SKIPPED,
            ],
        )

    @patch("household_validation.services.GroupIndividual.objects.select_for_update")
    def test_replace_primary_worker_clears_old_worker_before_assigning_new_one(
        self,
        select_for_update_mock,
    ):
        previous_worker = SimpleNamespace(
            individual_id="member-1",
            json_ext={"primary_worker": True},
        )
        selected_worker = SimpleNamespace(
            individual_id="member-2",
            json_ext={"primary_worker": False},
        )
        other_member = SimpleNamespace(
            individual_id="member-3",
            json_ext={},
        )
        select_for_update_mock.return_value.filter.return_value = [
            previous_worker,
            selected_worker,
            other_member,
        ]
        rows = [
            self._uploaded_row(2, "member-1", None),
            self._uploaded_row(3, "member-2", True),
        ]
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=SimpleNamespace(id="group-1")),
            patch.object(service, "_apply_primary_worker") as apply_worker_mock,
        ):
            changed = service._replace_primary_worker(rows)

        self.assertTrue(changed)
        self.assertEqual(
            apply_worker_mock.call_args_list,
            [
                call(previous_worker, False),
                call(selected_worker, True),
            ],
        )

    def test_unchanged_household_preserves_existing_verification_date(self):
        group = SimpleNamespace(
            json_ext={
                "validation_status": "VERIFIED",
                "last_verified_date": "2026-09-10",
                "validation_project_id": None,
                "validation_project_name": None,
                "validation_project_selection_type": PROJECT_SELECTION_TYPE_INTENT,
                "validation_notes": None,
            },
            save=MagicMock(),
        )
        uploaded_row = self._uploaded_row(2, "member-1", True)
        uploaded_row = replace(uploaded_row, verified=True)
        service = HouseholdValidationUploadService()

        changed = service._apply_group_validation_if_changed(
            group=group,
            uploaded_row=uploaded_row,
            project=None,
            upload_date=date(2026, 9, 11),
            uploaded_at=datetime(2026, 9, 11, 9, 0),
        )

        self.assertFalse(changed)
        self.assertEqual(group.json_ext["last_verified_date"], "2026-09-10")
        group.save.assert_not_called()

    def test_apply_row_does_not_update_participant_without_accepted_validation(self):
        group = SimpleNamespace(id="group-1")
        group_individual = SimpleNamespace(
            individual_id="member-1",
            json_ext={"primary_worker": False},
        )
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={
                "group_uuid": "group-1",
                "member_uuid": "member-1",
                "form_number": "FORM-001",
            },
            verified=None,
            primary_worker=True,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_group_individual", return_value=group_individual),
            patch.object(service, "_project", return_value=None),
            patch.object(service, "_structural_errors", return_value=[]),
            patch.object(service, "_apply_primary_worker") as apply_worker_mock,
            patch.object(service, "_save_batch_row") as save_row_mock,
        ):
            errors, participant_updated = service._apply_row(
                uploaded_row,
                batch=MagicMock(),
                upload_date=date.today(),
                uploaded_at=datetime.now(),
            )

        self.assertEqual(errors, [])
        self.assertFalse(participant_updated)
        apply_worker_mock.assert_not_called()
        self.assertEqual(
            save_row_mock.call_args.kwargs["status"],
            HouseholdValidationBatchRow.Status.SKIPPED,
        )

    def test_only_valid_rows_enable_household_participant_updates(self):
        verified_row = self._uploaded_row(2, "member-1", True)
        unverified_row = UploadedValidationRow(
            row_number=3,
            values={
                "group_uuid": "group-2",
                "member_uuid": "member-2",
                "form_number": "FORM-002",
            },
            verified=None,
            primary_worker=True,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )
        invalid_verified_row = UploadedValidationRow(
            row_number=4,
            values={
                "group_uuid": "group-3",
                "member_uuid": "member-3",
                "form_number": "FORM-003",
            },
            verified=True,
            primary_worker=True,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )
        service = HouseholdValidationUploadService()

        with patch.object(
            service,
            "_resolve_row",
            side_effect=[
                ([], MagicMock(), MagicMock(), None),
                ([], MagicMock(), MagicMock(), None),
                (["Row 4: member was not found in group"], None, None, None),
            ],
        ):
            accepted = service._accepted_validation_group_keys(
                [verified_row, unverified_row, invalid_verified_row]
            )

        self.assertEqual(accepted, {"group-1", "group-2"})

    @patch("household_validation.services.parse_validation_workbook")
    def test_upload_counts_only_national_id_updates(self, parse_workbook_mock):
        rows = [
            self._uploaded_row(2, "member-1", True),
            self._uploaded_row(3, "member-2", False),
            self._uploaded_row(4, "member-3", None),
        ]
        parse_workbook_mock.return_value = SimpleNamespace(
            rows=rows,
            rows_read=3,
            errors=[],
        )
        batch = MagicMock(id="batch-1")
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_accepted_validation_group_keys", return_value={"group-1"}),
            patch.object(service, "_primary_worker_rejections", return_value={}),
            patch.object(
                service,
                "_primary_worker_verification_statuses",
                return_value={"group-1": True},
            ),
            patch.object(service, "_replace_primary_worker"),
            patch.object(service, "_apply_group_validation_if_changed"),
            patch.object(service, "_get_or_create_batch", return_value=batch),
            patch.object(
                service,
                "_apply_row",
                side_effect=[([], False), ([], True), ([], False)],
            ),
            patch(
                "household_validation.services.transaction.atomic",
                return_value=nullcontext(),
            ),
        ):
            result = service.upload(b"workbook")

        self.assertEqual(result["participant_updates"], 1)
        self.assertEqual(result["households_verified"], 1)
        self.assertEqual(result["participants_verified"], 1)
        self.assertEqual(result["participants_not_verified"], 2)
        self.assertEqual(result["households_not_verified"], 0)

    @patch("household_validation.services.parse_validation_workbook")
    def test_parser_error_prevents_updates_for_the_whole_household(
        self,
        parse_workbook_mock,
    ):
        row = self._uploaded_row(2, "member-1", True)
        parse_workbook_mock.return_value = SimpleNamespace(
            rows=[row],
            rows_read=2,
            errors=["Row 3: primary_worker must be YES or NO"],
            invalid_group_keys=frozenset({"group-1"}),
        )
        batch = MagicMock(id="batch-1")
        service = HouseholdValidationUploadService()

        with (
            patch.object(
                service,
                "_accepted_validation_group_keys",
                return_value={"group-1"},
            ),
            patch.object(
                service,
                "_primary_worker_rejections",
                return_value={},
            ) as rejection_mock,
            patch.object(service, "_get_or_create_batch", return_value=batch),
            patch.object(service, "_replace_primary_worker") as replace_worker_mock,
            patch.object(service, "_apply_row", return_value=([], False)),
            patch.object(service, "_apply_group_validation_if_changed"),
            patch(
                "household_validation.services.transaction.atomic",
                return_value=nullcontext(),
            ),
        ):
            result = service.upload(b"workbook")

        self.assertEqual(
            rejection_mock.call_args.kwargs["eligible_group_keys"],
            set(),
        )
        replace_worker_mock.assert_not_called()
        self.assertEqual(result["rows_read"], 2)
        self.assertEqual(result["households_verified"], 0)
        self.assertEqual(result["participants_verified"], 0)
        self.assertEqual(result["errors"], 1)

    @patch("household_validation.services.parse_validation_workbook")
    def test_upload_counts_only_successful_participants_as_not_verified(
        self,
        parse_workbook_mock,
    ):
        rows = [
            self._uploaded_row(2, "member-1", False),
            self._uploaded_row(3, "member-2", False),
            self._uploaded_row(4, "member-3", None),
        ]
        parse_workbook_mock.return_value = SimpleNamespace(
            rows=rows,
            rows_read=3,
            errors=[],
        )
        batch = MagicMock(id="batch-1")
        service = HouseholdValidationUploadService()

        with (
            patch.object(service, "_accepted_validation_group_keys", return_value={"group-1"}),
            patch.object(service, "_primary_worker_rejections", return_value={}),
            patch.object(
                service,
                "_primary_worker_verification_statuses",
                return_value={"group-1": False},
            ),
            patch.object(service, "_replace_primary_worker") as replace_worker_mock,
            patch.object(service, "_apply_group_validation_if_changed"),
            patch.object(service, "_get_or_create_batch", return_value=batch),
            patch.object(
                service,
                "_apply_row",
                side_effect=[([], False), ([], False), (["invalid row"], False)],
            ),
            patch(
                "household_validation.services.transaction.atomic",
                return_value=nullcontext(),
            ),
        ):
            result = service.upload(b"workbook")

        self.assertEqual(result["households_verified"], 0)
        self.assertEqual(result["participants_verified"], 0)
        self.assertEqual(result["households_not_verified"], 1)
        self.assertEqual(result["participants_not_verified"], 2)
        self.assertEqual(result["participants_rejected"], 0)
        self.assertEqual(result["errors"], 1)
        replace_worker_mock.assert_not_called()

    @patch("household_validation.services.parse_validation_workbook")
    def test_upload_summary_uses_explicit_workbook_primary_workers(
        self,
        parse_workbook_mock,
    ):
        household_workers = {
            "group-1": [True],
            "group-2": [True, None],
            "group-3": [True, True, None],
            "group-4": [True, None],
            "group-5": [None, None],
            "group-6": [True],
            "group-7": [None, True],
            "group-8": [None, True],
        }
        rows = []
        row_number = 2
        for group_key, primary_workers in household_workers.items():
            for member_number, primary_worker in enumerate(
                primary_workers,
                start=1,
            ):
                row = self._uploaded_row(
                    row_number,
                    f"{group_key}-member-{member_number}",
                    primary_worker,
                )
                row.values["group_uuid"] = group_key
                rows.append(row)
                row_number += 1

        parse_workbook_mock.return_value = SimpleNamespace(
            rows=rows,
            rows_read=15,
            errors=[],
        )
        rejected = {
            "group-3": {
                "row_count": 3,
                "row_numbers": {5, 6, 7},
            }
        }
        statuses = {
            "group-1": True,
            "group-2": True,
            "group-4": True,
            "group-5": False,
            "group-6": True,
            "group-7": True,
            "group-8": True,
        }
        batch = MagicMock(id="batch-1")
        service = HouseholdValidationUploadService()

        with (
            patch.object(
                service,
                "_accepted_validation_group_keys",
                return_value=set(household_workers),
            ),
            patch.object(
                service,
                "_primary_worker_rejections",
                return_value=rejected,
            ),
            patch.object(
                service,
                "_primary_worker_verification_statuses",
                return_value=statuses,
            ),
            patch.object(service, "_replace_primary_worker") as replace_worker_mock,
            patch.object(service, "_apply_group_validation_if_changed"),
            patch.object(service, "_get_or_create_batch", return_value=batch),
            patch.object(service, "_save_primary_worker_rejection"),
            patch.object(service, "_apply_row", return_value=([], False)),
            patch(
                "household_validation.services.transaction.atomic",
                return_value=nullcontext(),
            ),
        ):
            result = service.upload(b"workbook")

        self.assertEqual(result["rows_read"], 15)
        self.assertEqual(result["participants_verified"], 6)
        self.assertEqual(result["participants_not_verified"], 6)
        self.assertEqual(result["participants_rejected"], 2)
        self.assertEqual(result["households_verified"], 6)
        self.assertEqual(result["households_not_verified"], 1)
        self.assertEqual(result["households_with_multiple_primary_workers"], 1)
        self.assertEqual(result["participant_updates"], 0)
        self.assertEqual(result["errors"], 0)
        self.assertEqual(replace_worker_mock.call_count, 6)

    @patch("household_validation.services.Group.objects.filter")
    def test_group_lookup_prefetches_members_and_individuals(self, filter_mock):
        group = SimpleNamespace(id="group-1")
        filtered_queryset = MagicMock()
        location_queryset = MagicMock()
        prefetched_queryset = MagicMock()
        filter_mock.return_value = filtered_queryset
        filtered_queryset.select_related.return_value = location_queryset
        location_queryset.prefetch_related.return_value = prefetched_queryset
        prefetched_queryset.first.return_value = group

        service = HouseholdValidationUploadService()
        result = service._group("group-1")
        cached_result = service._group("group-1")

        self.assertIs(result, group)
        self.assertIs(cached_result, group)
        filter_mock.assert_called_once_with(id="group-1")
        filtered_queryset.select_related.assert_called_once_with("location")
        location_queryset.prefetch_related.assert_called_once_with(
            "groupindividuals__individual"
        )

    def test_group_individual_lookup_reuses_prefetched_members(self):
        member = SimpleNamespace(
            individual_id="individual-1",
            is_deleted=False,
        )
        group = SimpleNamespace(
            _prefetched_objects_cache={"groupindividuals": [member]},
        )
        service = HouseholdValidationUploadService()

        with patch(
            "household_validation.services.GroupIndividual.objects.filter"
        ) as filter_mock:
            result = service._group_individual("individual-1", group)

        self.assertIs(result, member)
        filter_mock.assert_not_called()

    @patch("household_validation.services.GroupIndividual.objects.filter")
    def test_group_individual_lookup_uses_exported_individual_id(
        self,
        filter_mock,
    ):
        group = SimpleNamespace(id="group-1")
        group_individual = SimpleNamespace(
            id="membership-1",
            individual_id="individual-1",
            group=group,
        )
        queryset = MagicMock()
        queryset.select_related.return_value.first.return_value = group_individual
        filter_mock.return_value = queryset

        result = HouseholdValidationUploadService()._group_individual(
            "individual-1",
            group=group,
        )

        self.assertIs(result, group_individual)
        filter_mock.assert_called_once_with(
            individual_id="individual-1",
            group=group,
            is_deleted=False,
        )
        queryset.select_related.assert_called_once_with("individual")

    def test_json_safe_converts_nested_dates_to_iso_strings(self):
        raw_row = {
            "member_dob": date(1963, 10, 21),
            "validation": {
                "validation_date": date(2026, 7, 29),
            },
        }

        serialized = _json_safe(raw_row)

        self.assertEqual(serialized["member_dob"], "1963-10-21")
        self.assertEqual(
            serialized["validation"]["validation_date"],
            "2026-07-29",
        )

    def test_build_validation_json_ext_persists_not_verified_status(self):
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={},
            verified=False,
            primary_worker=None,
            validation_date=date(2026, 7, 3),
            project_name="Road Works",
            project_id="project-1",
            notes="Not available",
        )

        json_ext = build_validation_json_ext(
            uploaded_row=uploaded_row,
            project=None,
            upload_date=date(2026, 7, 4),
            uploaded_at=SimpleNamespace(isoformat=lambda: "2026-07-04T09:00:00+00:00"),
            user_id="user-1",
        )

        self.assertEqual(json_ext["validation_status"], VALIDATION_STATUS_NOT_VERIFIED)
        self.assertEqual(json_ext["last_verified_date"], "2026-07-03")
        self.assertEqual(json_ext["validation_project_selection_type"], PROJECT_SELECTION_TYPE_INTENT)

    def test_member_structural_errors_detect_protected_field_tampering(self):
        group = SimpleNamespace(
            json_ext={
                "form_number": "FORM-001",
                "household_wealth_quintile": "Poorest",
            },
        )
        group_individual = SimpleNamespace(
            role="HEAD",
            recipient_type="PRIMARY",
            group=group,
            individual=SimpleNamespace(
                first_name="Ada",
                last_name="Worker",
                dob=date(1990, 1, 1),
                json_ext={
                    "gender": "Female",
                    "fit_for_work": True,
                    "national_id": "NAT-001",
                },
            ),
        )
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={
                "form_number": "FORM-999",
                "member_name": "Grace Worker",
                "national_id": "NAT-999",
                "member_gender": "Male",
                "member_dob": "1991-01-01",
                "member_age": "18",
                "fit_for_work": "NO",
                "relationship": "SPOUSE",
                "head": "NO",
                "household_wealth_quintile": "Richest",
            },
            verified=None,
            primary_worker=None,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )

        errors = member_structural_errors(uploaded_row, group_individual)

        self.assertIn("Row 2: form_number does not match the household member", errors)
        self.assertIn("Row 2: member_name does not match the household member", errors)
        self.assertNotIn("Row 2: national_id does not match the household member", errors)
        self.assertIn("Row 2: member_gender does not match the household member", errors)
        self.assertIn("Row 2: fit_for_work does not match the household member", errors)
        self.assertIn("Row 2: relationship does not match the household member", errors)
        self.assertIn(
            "Row 2: household_wealth_quintile does not match the household member",
            errors,
        )

    def test_structural_errors_accept_integer_like_excel_ids(self):
        group = SimpleNamespace(
            json_ext={
                "form_number": "123456",
                "household_wealth_quintile": "Poorest",
            },
        )
        group_individual = SimpleNamespace(
            role="HEAD",
            recipient_type=None,
            group=group,
            individual=SimpleNamespace(
                first_name="Ada",
                last_name="Worker",
                dob=None,
                json_ext={
                    "fit_for_work": True,
                    "national_id": "123456",
                },
            ),
        )
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={
                "form_number": 123456.0,
                "national_id": 123456.0,
                "relationship": "HEAD",
                "household_wealth_quintile": "Poorest",
            },
            verified=None,
            primary_worker=None,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )

        errors = member_structural_errors(uploaded_row, group_individual)

        self.assertNotIn(
            "Row 2: form_number does not match the household member",
            errors,
        )
        self.assertNotIn(
            "Row 2: national_id does not match the household member",
            errors,
        )

    def test_form_number_maps_to_group_without_using_group_code(self):
        group = SimpleNamespace(
            code="INTERNAL-GROUP-CODE",
            json_ext={},
        )
        group_individual = SimpleNamespace(
            role="HEAD",
            recipient_type=None,
            group=group,
            individual=SimpleNamespace(
                first_name="Ada",
                last_name="Worker",
                dob=None,
                json_ext={
                    "fit_for_work": True,
                    "form_number": "FORM-001",
                    "national_id": "NAT-001",
                },
            ),
        )
        group.groupindividuals = _FakeRelatedManager([group_individual])
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={
                "form_number": "FORM-001",
                "national_id": "NAT-001",
                "relationship": "HEAD",
            },
            verified=None,
            primary_worker=None,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )

        errors = member_structural_errors(
            uploaded_row,
            group_individual,
            group=group,
        )

        self.assertNotIn(
            "Row 2: form_number does not match the household member",
            errors,
        )
        self.assertEqual(group.code, "INTERNAL-GROUP-CODE")

    def test_upload_wealth_validation_uses_head_before_first_member(self):
        other_member = self._wealth_member(
            "other",
            role="SPOUSE",
            fit_for_work=True,
            wealth_quintile="Richest",
        )
        head = self._wealth_member(
            "head",
            role="HEAD",
            fit_for_work=True,
            wealth_quintile="Poorer",
        )
        group = SimpleNamespace(
            json_ext={},
            groupindividuals=_FakeRelatedManager([other_member, head]),
        )
        other_member.group = group
        uploaded_row = UploadedValidationRow(
            row_number=2,
            values={
                "relationship": "SPOUSE",
                "household_wealth_quintile": "Poorer",
            },
            verified=None,
            primary_worker=None,
            validation_date=None,
            project_name=None,
            project_id=None,
            notes=None,
        )

        errors = member_structural_errors(uploaded_row, other_member)

        self.assertNotIn(
            "Row 2: household_wealth_quintile does not match "
            "the household member",
            errors,
        )

    def test_wealth_resolver_ignores_members_who_are_not_fit_for_work(self):
        ineligible_member = self._wealth_member(
            "ineligible",
            role="SPOUSE",
            fit_for_work=False,
            wealth_quintile="Richest",
        )
        eligible_member = self._wealth_member(
            "eligible",
            role="SON",
            fit_for_work=True,
            wealth_quintile="Middle",
        )
        group = SimpleNamespace(
            json_ext={},
            groupindividuals=_FakeRelatedManager(
                [ineligible_member, eligible_member]
            ),
        )

        self.assertEqual(get_household_wealth_quintile(group), "Middle")

    def test_pmt_score_resolver_ignores_members_who_are_not_fit_for_work(self):
        ineligible_member = self._wealth_member(
            "ineligible",
            role="SPOUSE",
            fit_for_work=False,
            wealth_quintile="Richest",
            pmt_score=9.9,
        )
        eligible_member = self._wealth_member(
            "eligible",
            role="SON",
            fit_for_work=True,
            wealth_quintile="Middle",
            pmt_score=-1.5,
        )
        group = SimpleNamespace(
            json_ext={},
            groupindividuals=_FakeRelatedManager(
                [ineligible_member, eligible_member]
            ),
        )

        self.assertEqual(get_household_pmt_score(group), -1.5)

    def test_export_selection_uses_shared_wealth_precedence(self):
        other_member = self._wealth_member(
            "other",
            role="SPOUSE",
            fit_for_work=True,
            wealth_quintile="Richest",
            pmt_score=9.9,
        )
        head = self._wealth_member(
            "head",
            role="HEAD",
            fit_for_work=True,
            wealth_quintile="Poorest",
            pmt_score=-2.1,
        )
        group = SimpleNamespace(
            id="group-1",
            code="HH-001",
            json_ext={},
            groupindividuals=_FakeRelatedManager([other_member, head]),
        )

        household = EligibleHouseholdSelectionService()._build_household(group)

        self.assertEqual(household.wealth_quintile, "Poorest")
        self.assertEqual(household.pmt_score, -2.1)

    def _wealth_member(self, member_id, role, fit_for_work, wealth_quintile, pmt_score=None):
        json_ext = {
            "fit_for_work": fit_for_work,
            "household_wealth_quintile": wealth_quintile,
        }
        if pmt_score is not None:
            json_ext["household_pmt_score"] = pmt_score
        individual = SimpleNamespace(
            id=f"individual-{member_id}",
            dob=date(1990, 1, 1),
            json_ext=json_ext,
        )
        return SimpleNamespace(
            id=f"membership-{member_id}",
            individual_id=individual.id,
            individual=individual,
            role=role,
            recipient_type=None,
        )

    def test_build_validation_error_report_writes_error_rows_csv(self):
        batch = SimpleNamespace(
            id="batch-1",
            rows=_FakeRows(
                [
                    SimpleNamespace(
                        row_number=2,
                        status="ERROR",
                        raw_row={
                            "form_number": "FORM-001",
                            "group_uuid": "group-1",
                            "member_uuid": "member-1",
                        },
                        error_message="form_number does not match",
                    )
                ]
            ),
        )

        report = build_validation_error_report_csv(batch.id, batch.rows.order_by("row_number"))

        self.assertIn("batch_id,row_number,status,form_number,group_uuid,member_uuid,error_message", report)
        self.assertIn("batch-1,2,ERROR,FORM-001,group-1,member-1,form_number does not match", report)

    def test_validation_error_report_excludes_legacy_rejection_rows(self):
        system_error = SimpleNamespace(
            row_number=2,
            status="ERROR",
            raw_row={
                "form_number": "FORM-001",
                "group_uuid": "group-1",
                "member_uuid": "member-1",
            },
            json_ext={},
            error_message="form_number does not match",
        )
        rejection = SimpleNamespace(
            row_number=3,
            status="ERROR",
            raw_row={
                "form_number": "FORM-002",
                "group_uuid": "group-2",
                "member_uuid": "member-2",
            },
            json_ext={"error_code": "MULTIPLE_PRIMARY_WORKERS"},
            error_message="household has more than one primary worker",
        )
        batch = SimpleNamespace(
            id="batch-1",
            rows=_FakeRows([system_error, rejection]),
        )

        report = build_validation_error_report(batch)

        self.assertIn("form_number does not match", report)
        self.assertNotIn("household has more than one primary worker", report)

    def test_error_report_uses_legacy_group_code_as_form_number(self):
        row = SimpleNamespace(
            row_number=2,
            status="ERROR",
            raw_row={
                "group_code": "LEGACY-001",
                "group_uuid": "group-1",
                "member_uuid": "member-1",
            },
            error_message="legacy upload error",
        )

        report = build_validation_error_report_csv("batch-1", [row])

        self.assertIn(
            "batch-1,2,ERROR,LEGACY-001,group-1,member-1,legacy upload error",
            report,
        )


class BusinessVerificationRulesTest(TestCase):
    def setUp(self):
        patcher = patch.object(HouseholdValidationConfig, "business_columns_enabled", True)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_non_primary_business_answers_block_the_whole_household(self):
        fixture = ExcelValidationListExporterTest()
        for worker in ("NO", None):
            for column, value in ((HAS_BUSINESS_COLUMN, "Yes"), (HAS_BUSINESS_COLUMN, "No"),
                                  (BUSINESS_TYPE_COLUMN, "Crop farming"), (BUSINESS_DURATION_COLUMN, 0)):
                with self.subTest(worker=worker, column=column, value=value):
                    workbook = ExcelValidationListExporter(fixture._selection_result(member_count=2), "batch-1").export_workbook()
                    sheet = workbook[VALIDATION_LIST_SHEET]
                    fixture._cell(sheet, "primary_worker", 2).value = "YES"
                    fixture._cell(sheet, HAS_BUSINESS_COLUMN, 2).value = "No"
                    fixture._cell(sheet, "primary_worker", 3).value = worker
                    fixture._cell(sheet, column, 3).value = value
                    content = ValidationUploadParserTest()._workbook_bytes(workbook)
                    parsed = parse_validation_workbook(content)
                    self.assertIn("group-1", parsed.invalid_group_keys)
                    self.assertTrue(any("Clear all three business fields" in error for error in parsed.errors))
                    service = HouseholdValidationUploadService()
                    with patch.object(service, "_resolve_row", return_value=([], None, None, None)):
                        result = service.upload(content, dry_run=True)
                    self.assertGreater(result["errors"], 0)
                    self.assertEqual(result["households_verified"], 0)
                    self.assertEqual(result["participants_verified"], 0)

    def test_all_business_fields_have_primary_worker_gates_and_stale_value_highlights(self):
        fixture = ExcelValidationListExporterTest()
        workbook = ExcelValidationListExporter(fixture._selection_result(), "batch-1").export_workbook()
        sheet = workbook[VALIDATION_LIST_SHEET]
        for column in BUSINESS_COLUMNS:
            cell = fixture._cell(sheet, column)
            rule = next(rule for rule in sheet.data_validations.dataValidation if cell.coordinate in rule.sqref)
            self.assertIn('$P2', rule.formula1)
            self.assertLessEqual(len(rule.formula1), 255 if rule.type == "list" else 8192)
            self.assertEqual(rule.errorStyle, "stop")
            self.assertTrue(rule.showErrorMessage)
            self.assertFalse(rule.allow_blank)
            rules = [r for target in sheet.conditional_formatting if cell.coordinate in target.sqref
                     for r in sheet.conditional_formatting[target]]
            self.assertTrue(any('NOT(' in r.formula[0] and 'LEN(TRIM(' in r.formula[0] for r in rules))
            self.assertFalse(any(r.dxf.fill.fgColor.rgb == 'FFE7E6E6' for r in rules))
            self.assertEqual(cell.fill.fgColor.rgb, fixture._cell(sheet, "validation_notes").fill.fgColor.rgb)

    def test_business_rejections_are_included_in_download(self):
        from household_validation.verification import BUSINESS_REJECTION_CODE

        rows = [SimpleNamespace(
            json_ext={"error_code": BUSINESS_REJECTION_CODE},
            error_message="business owner is not a primary worker",
            group_id="group-1", row_number=number,
            raw_row={"form_number": "FORM-001", "group_uuid": "group-1"},
        ) for number in (2, 3)]
        content, count = build_rejected_households_workbook_bytes(rows)
        sheet = load_workbook(BytesIO(content)).active
        self.assertEqual(count, 1)
        self.assertEqual(sheet.cell(2, 1).value, "FORM-001")
        self.assertEqual(sheet.cell(2, 3).value, "2, 3")
        self.assertEqual(sheet.cell(2, 4).value, rows[0].error_message)

    def test_completed_workbook_rules_match_dry_run_and_persistence(self):
        from household_validation.verification import VERIFIED, NOT_VERIFIED, REJECTED

        scenarios = (
            ("YES", "Yes", VERIFIED),
            ("YES", "No", VERIFIED),
            ("YES", None, NOT_VERIFIED),
            ("NO", None, NOT_VERIFIED),
            (None, None, NOT_VERIFIED),
        )
        for worker, business, expected in scenarios:
            with self.subTest(worker=worker, business=business):
                workbook = ValidationUploadParserTest()._upload_workbook()
                sheet = workbook[VALIDATION_LIST_SHEET]
                for column, value in (("primary_worker", worker), (HAS_BUSINESS_COLUMN, business)):
                    sheet.cell(2, EXCEL_COLUMNS.index(column) + 1, value)
                if business == "Yes":
                    sheet.cell(2, EXCEL_COLUMNS.index(BUSINESS_TYPE_COLUMN) + 1, "Crop farming")
                    sheet.cell(2, EXCEL_COLUMNS.index(BUSINESS_DURATION_COLUMN) + 1, 1)
                # Never trust status text or cached formula values supplied by the client.
                for column in ("participant_status", "household_status"):
                    sheet.cell(2, EXCEL_COLUMNS.index(column) + 1, "FORGED")
                content = ValidationUploadParserTest()._workbook_bytes(workbook)
                individual = SimpleNamespace(id="member-1", json_ext={})
                member = SimpleNamespace(
                    individual_id="member-1", individual=individual, is_deleted=False,
                )
                group = SimpleNamespace(
                    id="group-1", json_ext={}, save=MagicMock(),
                    _prefetched_objects_cache={"groupindividuals": [member]},
                )
                service = HouseholdValidationUploadService()
                with (
                    patch.object(service, "_group", return_value=group),
                    patch.object(service, "_structural_errors", return_value=[]),
                    patch.object(service, "_get_or_create_batch", return_value=MagicMock()) as batch,
                    patch.object(service, "_replace_primary_worker", return_value=False) as replace_worker,
                    patch.object(service, "_save_batch_row") as audit,
                    patch("household_validation.services.IndividualService") as individual_service,
                    patch("household_validation.services.transaction.atomic", return_value=nullcontext()),
                ):
                    dry_run = service.upload(content, dry_run=True)
                    batch.assert_not_called()
                    individual_service.assert_not_called()
                    result = service.upload(content)
                self.assertEqual(result["errors"], 0)
                for status, suffix in ((VERIFIED, "verified"), (NOT_VERIFIED, "not_verified"), (REJECTED, "rejected")):
                    for prefix in ("participants", "households"):
                        key = f"{prefix}_{suffix}"
                        self.assertEqual(result[key], int(expected == status))
                        self.assertEqual(dry_run[key], result[key])
                self.assertEqual(individual.json_ext["validation_status"], expected)
                self.assertEqual(group.json_ext["validation_status"], expected)
                self.assertEqual(replace_worker.call_count, int(expected == VERIFIED))
                self.assertEqual(
                    audit.call_args.kwargs["status"],
                    HouseholdValidationBatchRow.Status.REJECTED if expected == REJECTED
                    else HouseholdValidationBatchRow.Status.APPLIED,
                )

    def test_rejection_takes_priority_in_household_with_mixed_members(self):
        from household_validation.verification import household_status

        self.assertEqual(household_status([(True, "Yes"), (False, None)]), "VERIFIED")
        self.assertEqual(household_status([(True, "No"), (None, None)]), "VERIFIED")
        self.assertEqual(household_status([(True, "Yes"), (False, "No")]), "VERIFIED")
        self.assertEqual(household_status([(True, "Yes"), (True, "No")]), "REJECTED")

    def test_export_has_protected_calculated_status_columns(self):
        workbook = ExcelValidationListExporter(
            ExcelValidationListExporterTest()._selection_result(), "batch-1",
        ).export_workbook()
        worksheet = workbook[VALIDATION_LIST_SHEET]
        headers = [cell.value for cell in worksheet[1]]
        for column in ("participant_status", "household_status"):
            cell = worksheet.cell(2, headers.index(column) + 1)
            self.assertEqual(cell.data_type, "f")
            self.assertTrue(cell.protection.locked)
        self.assertEqual(workbook.calculation.calcMode, "auto")


class PwpVerificationRulesTest(TestCase):
    def setUp(self):
        patcher = patch.object(HouseholdValidationConfig, "business_columns_enabled", False)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_pwp_verification_depends_only_on_primary_worker(self):
        from household_validation.verification import participant_status, household_status

        for business in (None, "No", "Yes"):
            for worker, expected in ((True, "VERIFIED"), (False, "NOT_VERIFIED"), (None, "NOT_VERIFIED")):
                with self.subTest(worker=worker, business=business):
                    self.assertEqual(participant_status(
                        worker, business, business_columns_enabled=False,
                    ), expected)
                    self.assertEqual(household_status(
                        [(worker, business), (False, "Yes")], business_columns_enabled=False,
                    ), expected)
        self.assertEqual(household_status(
            [(True, None), (True, None)], business_columns_enabled=False,
        ), "REJECTED")

    def test_pwp_generated_workbook_upload_matches_dry_run_and_saved_statuses(self):
        fixture = ExcelValidationListExporterTest()
        for worker in ("YES", " yes ", "Y", "TRUE", "1", "NO", None):
            with self.subTest(worker=worker):
                verified = worker not in ("NO", None)
                expected = "VERIFIED" if verified else "NOT_VERIFIED"
                workbook = ExcelValidationListExporter(
                    fixture._selection_result(member_count=2), "batch-1",
                ).export_workbook()
                sheet = workbook[VALIDATION_LIST_SHEET]
                fixture._cell(sheet, "primary_worker").value = worker
                # Upload must recompute decisions, not trust status cells.
                for row in (2, 3):
                    for column in ("participant_status", "household_status"):
                        fixture._cell(sheet, column, row).value = "FORGED"
                content = ValidationUploadParserTest()._workbook_bytes(workbook)
                individuals = [SimpleNamespace(
                    id=f"member-{n}", json_ext={"business_experience": "Yes", "business_period": 2},
                ) for n in (1, 2)]
                members = [SimpleNamespace(
                    individual_id=individual.id, individual=individual, is_deleted=False,
                ) for individual in individuals]
                group = SimpleNamespace(
                    id="group-1", json_ext={}, save=MagicMock(),
                    _prefetched_objects_cache={"groupindividuals": members},
                )
                service = HouseholdValidationUploadService()
                with (
                    patch.object(service, "_group", return_value=group),
                    patch.object(service, "_structural_errors", return_value=[]),
                    patch.object(service, "_get_or_create_batch", return_value=MagicMock()) as batch,
                    patch.object(service, "_replace_primary_worker", return_value=False) as replace_worker,
                    patch.object(service, "_save_batch_row"),
                    patch("household_validation.services.IndividualService") as individual_service,
                    patch("household_validation.services.transaction.atomic", return_value=nullcontext()),
                ):
                    dry_run = service.upload(content, dry_run=True)
                    batch.assert_not_called()
                    individual_service.assert_not_called()
                    result = service.upload(content)
                self.assertEqual(result["errors"], 0, result["error_messages"])
                for key, count in {
                    "participants_verified": int(verified),
                    "participants_not_verified": 2 - int(verified),
                    "participants_rejected": 0,
                    "households_verified": int(verified),
                    "households_not_verified": int(not verified),
                    "households_rejected": 0,
                }.items():
                    self.assertEqual(result[key], count)
                    self.assertEqual(dry_run[key], count)
                self.assertEqual(individuals[0].json_ext["validation_status"], expected)
                self.assertEqual(individuals[1].json_ext["validation_status"], "NOT_VERIFIED")
                self.assertEqual(group.json_ext["validation_status"], expected)
                self.assertEqual(replace_worker.call_count, int(verified))
                for individual in individuals:
                    self.assertEqual(individual.json_ext["business_experience"], "Yes")
                    self.assertEqual(individual.json_ext["business_period"], 2)

    def test_pwp_upload_rejects_multiple_primary_workers_without_updates(self):
        fixture = ExcelValidationListExporterTest()
        workbook = ExcelValidationListExporter(
            fixture._selection_result(member_count=2), "batch-1",
        ).export_workbook()
        sheet = workbook[VALIDATION_LIST_SHEET]
        for row in (2, 3):
            fixture._cell(sheet, "primary_worker", row).value = "YES"
        content = ValidationUploadParserTest()._workbook_bytes(workbook)
        members = [SimpleNamespace(
            individual_id=f"member-{n}", is_deleted=False,
            individual=SimpleNamespace(id=f"member-{n}", json_ext={}),
        ) for n in (1, 2)]
        group = SimpleNamespace(
            id="group-1", json_ext={},
            _prefetched_objects_cache={"groupindividuals": members},
        )
        service = HouseholdValidationUploadService()
        with (
            patch.object(service, "_group", return_value=group),
            patch.object(service, "_structural_errors", return_value=[]),
            patch.object(service, "_get_or_create_batch", return_value=MagicMock()),
            patch.object(service, "_save_primary_worker_rejection"),
            patch.object(service, "_apply_row") as apply_row,
            patch.object(service, "_replace_primary_worker") as replace_worker,
            patch("household_validation.services.transaction.atomic", return_value=nullcontext()),
        ):
            dry_run = service.upload(content, dry_run=True)
            result = service.upload(content)
        apply_row.assert_not_called()
        replace_worker.assert_not_called()
        for summary in (dry_run, result):
            self.assertEqual(summary["households_rejected"], 1)
            self.assertEqual(summary["participants_rejected"], 2)
            self.assertEqual(summary["households_verified"], 0)


class LocalDateTests(TestCase):
    """openIMIS runs with USE_TZ = False, so timezone.now() is naive.

    Passing a naive datetime to timezone.localdate()/localtime() raises
    ValueError("localtime() cannot be applied to a naive datetime"), which
    previously aborted every non-dry-run validation-list upload.
    """

    def test_naive_datetime_returns_its_date(self):
        naive = datetime(2026, 7, 29, 12, 27)
        self.assertIsNone(naive.tzinfo)
        self.assertEqual(_local_date(naive), date(2026, 7, 29))

    def test_aware_datetime_is_converted_to_local_date(self):
        aware = datetime(2026, 7, 29, 12, 27, tzinfo=datetime_timezone.utc)
        self.assertEqual(_local_date(aware), timezone.localtime(aware).date())


class _FakeRows:
    def __init__(self, rows):
        self.rows = rows

    def filter(self, **kwargs):
        return self

    def order_by(self, *fields):
        return self.rows
