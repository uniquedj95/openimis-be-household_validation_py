from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from household_validation.identity import get_household_form_number
from household_validation.verification import (
    PARTICIPANT_STATUS_COLUMN, HOUSEHOLD_STATUS_COLUMN, BUSINESS_REJECTION_CODE,
)


PRIMARY_WORKER_ANSWERS_RANGE = "HouseholdPrimaryWorkerAnswers"
PRIMARY_WORKER_NO_RANGE = "HouseholdPrimaryWorkerNo"
BUSINESS_ANSWERS_RANGE = "HouseholdBusinessAnswers"
BUSINESS_UNAVAILABLE_RANGE = "HouseholdBusinessUnavailable"
HOUSEHOLD_ROW_COLORS = ("FFA9D18E", "FFE2F0D9")

LOCATION_COLUMN_TYPES = {
    "District": "R",
    "TA": "D",
    "GVH": "W",
    "Village": "V",
}

MICRO_CATCHMENT_COLUMN = "Micro-Catchment"
HOTSPOT_COLUMN = "Hotspot"

HAS_BUSINESS_COLUMN = "Does member has a business"
BUSINESS_TYPE_COLUMN = "Type of Business"
BUSINESS_DURATION_COLUMN = "Business Period (in years)"
BUSINESS_COLUMNS = {HAS_BUSINESS_COLUMN, BUSINESS_TYPE_COLUMN, BUSINESS_DURATION_COLUMN}

BUSINESS_TYPE_OPTIONS_SHEET = "Business Type Options"

# Full upload schema, including optional fields from previously issued workbooks.
# Each exporter selects its output columns from this schema using deployment config.
EXCEL_COLUMNS = [
    "batch_id",
    "group_uuid",
    "member_uuid",
    "row_type",
    "District",
    MICRO_CATCHMENT_COLUMN,
    "TA",
    "GVH",
    HOTSPOT_COLUMN,
    "Village",
    "form_number",
    "member_name",
    "relationship",
    "member_dob",
    "national_id",
    "primary_worker",
    "member_gender",
    "member_age",
    "marital_status",
    "disability",
    "fit_for_work",
    "pmt_score",
    "household_wealth_quintile",
    "project",
    "project_id",
    HAS_BUSINESS_COLUMN,
    BUSINESS_TYPE_COLUMN,
    BUSINESS_DURATION_COLUMN,
    "validation_notes",
    PARTICIPANT_STATUS_COLUMN,
    HOUSEHOLD_STATUS_COLUMN,
]

PROJECT_OPTIONS_SHEET = "Project Options"
PROJECT_OPTIONS_HEADERS = ["project_id", "project", "project_label"]

EDITABLE_COLUMNS = {
    "national_id",
    "primary_worker",
    "project",
    HAS_BUSINESS_COLUMN,
    BUSINESS_TYPE_COLUMN,
    BUSINESS_DURATION_COLUMN,
    "validation_notes",
}

TEXT_COLUMNS = {
    "form_number",
    "national_id",
}

PRIMARY_WORKER_REJECTION_CODE = "MULTIPLE_PRIMARY_WORKERS"
PRIMARY_WORKER_REJECTION_MESSAGE = (
    "household has more than one primary worker"
)
LEGACY_PRIMARY_WORKER_REJECTION_MESSAGE = (
    "Rejected: household would have more than one primary worker"
)


def is_primary_worker_rejection(row):
    json_ext = row.json_ext or {}
    return (
        json_ext.get("error_code") == PRIMARY_WORKER_REJECTION_CODE
        or row.error_message
        in {
            PRIMARY_WORKER_REJECTION_MESSAGE,
            LEGACY_PRIMARY_WORKER_REJECTION_MESSAGE,
        }
    )


def build_rejected_households_workbook_bytes(rows):
    rejected_rows = [
        row for row in rows
        if is_primary_worker_rejection(row)
        or (row.json_ext or {}).get("error_code") == BUSINESS_REJECTION_CODE
    ]
    households = {}
    for row in rejected_rows:
        raw_row = row.raw_row or {}
        group_uuid = str(
            getattr(row, "group_id", None)
            or raw_row.get("group_uuid")
            or f"row-{row.row_number}"
        )
        household = households.setdefault(
            group_uuid,
            {
                "form_number": (
                    raw_row.get("form_number")
                    or raw_row.get("group_code")
                    or group_uuid
                ),
                "group_uuid": group_uuid,
                "row_numbers": set(),
                "rejection_reason": row.error_message,
            },
        )
        if row.row_number is not None:
            household["row_numbers"].add(row.row_number)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Rejected Households"
    headers = [
        "form_number",
        "group_uuid",
        "workbook_rows",
        "rejection_reason",
    ]
    worksheet.append(headers)

    for household in households.values():
        values = [
            household["form_number"],
            household["group_uuid"],
            ", ".join(
                str(row_number)
                for row_number in sorted(household["row_numbers"])
            ),
            household["rejection_reason"],
        ]
        worksheet.append(values)
        for cell in worksheet[worksheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, column_cells in enumerate(worksheet.columns, start=1):
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_length + 2, 12),
            50,
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), len(households)

def _configured_business_type_options():
    from household_validation.apps import (
        DEFAULT_BUSINESS_TYPE_OPTIONS,
        HouseholdValidationConfig,
    )

    value = getattr(HouseholdValidationConfig, "business_type_options", None)
    if not isinstance(value, (list, tuple)):
        return list(DEFAULT_BUSINESS_TYPE_OPTIONS)
    options = [str(option).strip() for option in value if str(option).strip()]
    return options or list(DEFAULT_BUSINESS_TYPE_OPTIONS)


class ExcelValidationListExporter:
    def __init__(self, selection_result, batch_id, projects=None):
        from household_validation.apps import HouseholdValidationConfig

        self.selection_result = selection_result
        self.batch_id = batch_id
        self.projects = projects or []
        self.business_columns_enabled = HouseholdValidationConfig.business_columns_enabled
        self.columns = [
            column for column in EXCEL_COLUMNS
            if self.business_columns_enabled or column not in BUSINESS_COLUMNS
        ]
        self.business_type_options = _configured_business_type_options()
        self._micro_catchment_cache = {}
        self._hotspot_cache = {}

    def export_workbook(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Validation List"
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True

        self._write_header(worksheet)
        self._write_rows(worksheet)
        self._write_status_formulas(worksheet)
        project_options_worksheet = self._write_project_options(workbook)
        if self.business_columns_enabled:
            business_type_options_worksheet = self._write_business_type_options(workbook)
            self._write_primary_worker_options(workbook, business_type_options_worksheet)
            self._autosize_columns(business_type_options_worksheet)
        else:
            self._write_primary_worker_options(workbook, project_options_worksheet)
        self._apply_validation(worksheet)
        self._apply_protection(worksheet)
        self._autosize_columns(worksheet)
        self._autosize_columns(project_options_worksheet)

        # Keep identifiers in each row for reliable upload matching and formulas,
        # but omit these internal fields from the field officer's visible sheet.
        for column in ("batch_id", "group_uuid", "member_uuid", "row_type", "project_id"):
            worksheet.column_dimensions[self._column_letter(column)].hidden = True
        for selection in worksheet.sheet_view.selection:
            selection.activeCell = "E2"
            selection.sqref = "E2"
        return workbook

    def export_bytes(self):
        output = BytesIO()
        self.export_workbook().save(output)
        output.seek(0)
        return output.getvalue()

    def _write_header(self, worksheet):
        header_fill = PatternFill("solid", fgColor="D9EAD3")
        for column_number, title in enumerate(self.columns, start=1):
            cell = worksheet.cell(row=1, column=column_number, value=title)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.protection = Protection(locked=True)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

    def _write_rows(self, worksheet):
        household_fills = {}
        for row_number, selected_member in enumerate(self.selection_result.member_rows, start=2):
            household_key = str(selected_member.household.id)
            if household_key not in household_fills:
                color = HOUSEHOLD_ROW_COLORS[
                    len(household_fills) % len(HOUSEHOLD_ROW_COLORS)
                ]
                household_fills[household_key] = PatternFill(
                    fill_type="solid",
                    fgColor=color,
                )
            household_fill = household_fills[household_key]
            values = self._build_row(selected_member)
            for column_number, title in enumerate(self.columns, start=1):
                value = values.get(title)
                if title in TEXT_COLUMNS and value is not None:
                    value = str(value)
                cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )
                if title in TEXT_COLUMNS:
                    cell.number_format = "@"
                cell.fill = household_fill
                cell.protection = Protection(locked=title not in EDITABLE_COLUMNS)

    def _write_project_options(self, workbook):
        worksheet = workbook.create_sheet(PROJECT_OPTIONS_SHEET)
        for column_number, title in enumerate(PROJECT_OPTIONS_HEADERS, start=1):
            worksheet.cell(row=1, column=column_number, value=title)
        project_labels = self._project_labels()
        for row_number, project in enumerate(self.projects, start=2):
            worksheet.cell(row=row_number, column=1, value=self._project_id(project))
            worksheet.cell(row=row_number, column=2, value=self._project_name(project))
            worksheet.cell(row=row_number, column=3, value=project_labels[id(project)])
        worksheet.sheet_state = "hidden"
        return worksheet

    def _write_status_formulas(self, worksheet):
        worker = self._column_letter("primary_worker")
        business = self._column_letter(HAS_BUSINESS_COLUMN) if self.business_columns_enabled else None
        participant = self._column_letter(PARTICIPANT_STATUS_COLUMN)
        group = self._column_letter("group_uuid")
        last = worksheet.max_row
        groups = f'${group}$2:${group}${last}'
        workers = f'${worker}$2:${worker}${last}'
        statuses = f'${participant}$2:${participant}${last}'
        for row in range(2, last + 1):
            # Use the same accepted boolean spellings as the upload parser.
            def matches(column, options):
                cell = f'UPPER(TRIM({column}{row}&""))'
                return 'OR(' + ','.join(f'{cell}="{value}"' for value in options) + ')'
            primary_yes = matches(worker, ("YES", "Y", "TRUE", "1"))
            if self.business_columns_enabled:
                business_yes = matches(business, ("YES", "Y", "TRUE", "1"))
                business_no = matches(business, ("NO", "N", "FALSE", "0"))
                business_type = self._column_letter(BUSINESS_TYPE_COLUMN)
                period = self._column_letter(BUSINESS_DURATION_COLUMN)
                valid_details = (
                    f'AND(COUNTIF(\'{BUSINESS_TYPE_OPTIONS_SHEET}\'!$A$2:$A${len(self.business_type_options)+1},'
                    f'{business_type}{row})>0,ISNUMBER({period}{row}),'
                    f'{period}{row}>=0,{period}{row}<=100)'
                )
                participant_formula = (
                    f'=IF(AND({primary_yes},OR({business_no},AND({business_yes},{valid_details}))),'
                    f'"VERIFIED","NOT_VERIFIED")'
                )
            else:
                participant_formula = f'=IF({primary_yes},"VERIFIED","NOT_VERIFIED")'
            worksheet.cell(row, self.columns.index(PARTICIPANT_STATUS_COLUMN) + 1,
                participant_formula)
            worker_matches = '+'.join(
                f'(UPPER(TRIM({workers}&""))="{value}")'
                for value in ("YES", "Y", "TRUE", "1")
            )
            worker_count = f'SUMPRODUCT(--({groups}=${group}{row}),--(({worker_matches})>0))'
            worksheet.cell(row, self.columns.index(HOUSEHOLD_STATUS_COLUMN) + 1,
                f'=IF(OR({worker_count}>1,COUNTIFS({groups},${group}{row},{statuses},'
                f'"REJECTED")>0),"REJECTED",IF(COUNTIFS({groups},${group}{row},'
                f'{statuses},"VERIFIED")>0,"VERIFIED","NOT_VERIFIED"))')

    def _write_business_type_options(self, workbook):
        worksheet = workbook.create_sheet(BUSINESS_TYPE_OPTIONS_SHEET)
        worksheet.cell(row=1, column=1, value="business_type")
        for row_number, business_type in enumerate(self.business_type_options, start=2):
            worksheet.cell(row=row_number, column=1, value=business_type)
        # Named ranges keep the dependent dropdown compatible with Excel and Calc.
        worksheet["C1"] = "Business answers"
        worksheet["C2"] = "Yes"
        worksheet["C3"] = "No"
        # Calc displays a genuinely empty list-source cell as numeric 0.
        # An explicit empty string keeps the unavailable dropdown blank.
        worksheet["C4"] = '=""'
        for name, reference in (
            (BUSINESS_ANSWERS_RANGE, "$C$2:$C$3"),
            (BUSINESS_UNAVAILABLE_RANGE, "$C$4"),
        ):
            workbook.defined_names.add(DefinedName(
                name, attr_text=f"'{BUSINESS_TYPE_OPTIONS_SHEET}'!{reference}",
            ))
        worksheet.sheet_state = "hidden"
        return worksheet

    def _write_primary_worker_options(self, workbook, worksheet):
        # Primary Worker must still work when the business options sheet is absent.
        worksheet["E1"] = "Primary Worker answers"
        worksheet["E2"] = "YES"
        worksheet["E3"] = "NO"
        for name, reference in (
            (PRIMARY_WORKER_ANSWERS_RANGE, "$E$2:$E$3"),
            (PRIMARY_WORKER_NO_RANGE, "$E$3"),
        ):
            workbook.defined_names.add(DefinedName(
                name, attr_text=f"'{worksheet.title}'!{reference}",
            ))

    def _build_row(self, selected_member):
        household = selected_member.household
        member = selected_member.member
        group = household.source
        group_individual = member.source
        individual = getattr(group_individual, "individual", None)
        location = getattr(group, "location", None)

        return {
            "batch_id": str(self.batch_id),
            "row_type": selected_member.row_type,
            **{
                column: self._location_name(location, location_type)
                for column, location_type in LOCATION_COLUMN_TYPES.items()
            },
            MICRO_CATCHMENT_COLUMN: self._micro_catchment_name(location),
            HOTSPOT_COLUMN: self._hotspot_name(location),
            "form_number": get_household_form_number(group, individual),
            "group_uuid": str(household.id),
            "member_uuid": str(member.id),
            "member_name": self._member_name(individual),
            "national_id": self._national_id(individual),
            "member_gender": member.gender,
            "member_dob": member.dob,
            "member_age": member.age,
            "marital_status": self._marital_status(individual),
            "disability": self._disability(individual),
            "fit_for_work": "YES" if member.fit_for_work else "NO",
            "relationship": self._relationship(member.role),
            "pmt_score": household.pmt_score,
            "household_wealth_quintile": household.wealth_quintile,
            # Field-validation inputs start fresh on every export, even when
            # the individual already has stored answers from a previous visit.
            "primary_worker": None,
            "project": None,
            "project_id": None,
            HAS_BUSINESS_COLUMN: None,
            BUSINESS_TYPE_COLUMN: None,
            BUSINESS_DURATION_COLUMN: None,
            "validation_notes": None,
        }

    def _apply_validation(self, worksheet):
        max_row = max(worksheet.max_row, 2)
        primary_worker_col = self._column_letter("primary_worker")
        project_col = self._column_letter("project")

        group_col = self._column_letter("group_uuid")
        groups = f"${group_col}$2:${group_col}${max_row}"
        workers = f"${primary_worker_col}$2:${primary_worker_col}${max_row}"
        worker_count = f'COUNTIFS({groups},${group_col}2,{workers},"YES")'
        # Exclude this cell, so the selected worker can keep YES or change to NO.
        other_workers = f'{worker_count}-COUNTIF(${primary_worker_col}2,"YES")'
        primary_worker_validation = DataValidation(
            type="list",
            formula1=(
                f'INDIRECT(IF({other_workers}=0,'
                f'"{PRIMARY_WORKER_ANSWERS_RANGE}","{PRIMARY_WORKER_NO_RANGE}"))'
            ),
            allow_blank=True,
            showDropDown=False,
            showInputMessage=True,
            promptTitle="One Primary Worker only",
            prompt=(
                'Only one member per household can be YES. To change the '
                'Primary Worker, set the current worker to NO or clear it first.'
            ),
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Primary Worker already selected",
            error=(
                'This household already has a Primary Worker. Select NO, '
                "or clear the other member's YES before selecting this member."
            ),
        )
        worksheet.add_data_validation(primary_worker_validation)
        primary_worker_validation.add(
            f"{primary_worker_col}2:{primary_worker_col}{max_row}"
        )
        worksheet.conditional_formatting.add(
            f"{primary_worker_col}2:{primary_worker_col}{max_row}",
            FormulaRule(
                formula=[f'AND(${primary_worker_col}2="YES",{worker_count}>1)'],
                fill=PatternFill(fill_type="solid", fgColor="FFFFC7CE"),
                font=Font(color="FF9C0006"),
            ),
        )

        project_count = len([project for project in self.projects if self._project_name(project)])
        if project_count:
            project_formula = f"'{PROJECT_OPTIONS_SHEET}'!$C$2:$C${project_count + 1}"
            project_validation = DataValidation(
                type="list",
                formula1=project_formula,
                allow_blank=True,
            )
            worksheet.add_data_validation(project_validation)
            project_validation.add(f"{project_col}2:{project_col}{max_row}")

        if not self.business_columns_enabled:
            return

        has_business_col = self._column_letter(HAS_BUSINESS_COLUMN)
        business_type_col = self._column_letter(BUSINESS_TYPE_COLUMN)
        business_duration_col = self._column_letter(BUSINESS_DURATION_COLUMN)
        primary_worker_answered = (
            f'OR(TRIM(${primary_worker_col}2&"")="YES",'
            f'TRIM(${primary_worker_col}2&"")="Y",'
            f'TRIM(${primary_worker_col}2&"")="TRUE",'
            f'TRIM(${primary_worker_col}2&"")="1")'
        )
        has_business_validation = DataValidation(
            type="list",
            formula1=(
                f'INDIRECT(IF({primary_worker_answered},'
                f'"{BUSINESS_ANSWERS_RANGE}","{BUSINESS_UNAVAILABLE_RANGE}"))'
            ),
            # Ignore-blank must be off: otherwise an empty source can bypass
            # the prerequisite when a value is typed directly into the cell.
            allow_blank=False,
            showDropDown=False,
            showInputMessage=True,
            promptTitle="Select Primary Worker first",
            prompt=(
                'First select YES in Primary Worker on this row. '
                'Then choose Yes or No for Does member has a business. '
                'If Yes, Business Period is required.'
            ),
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Select Primary Worker first",
            error=(
                'Select YES in Primary Worker on this row first, '
                'then choose Yes or No from the Business dropdown.'
            ),
        )
        worksheet.add_data_validation(has_business_validation)
        has_business_validation.add(
            f"{has_business_col}2:{has_business_col}{max_row}"
        )

        business_type_options_range = (
            f"'{BUSINESS_TYPE_OPTIONS_SHEET}'!$A$2:$A${len(self.business_type_options) + 1}"
        )
        business_type_validation = DataValidation(
            type="list",
            formula1=f'IF(AND({primary_worker_answered},UPPER(TRIM(${has_business_col}2))="YES"),{business_type_options_range},{BUSINESS_UNAVAILABLE_RANGE})',
            allow_blank=False,
            errorStyle="stop",
            showInputMessage=True,
            promptTitle="Primary Worker YES required",
            prompt="Select Primary Worker YES and Business Yes before entering business details.",
        )
        business_type_validation.error = (
            'Select Primary Worker YES and Business Yes before choosing a business type.'
        )
        business_type_validation.errorTitle = "Business type not applicable"
        business_type_validation.showErrorMessage = True
        worksheet.add_data_validation(business_type_validation)
        business_type_validation.add(
            f"{business_type_col}2:{business_type_col}{max_row}"
        )

        business_yes = (
            f'AND({primary_worker_answered},OR(UPPER(TRIM(${has_business_col}2&""))="YES",'
            f'UPPER(TRIM(${has_business_col}2&""))="Y",'
            f'UPPER(TRIM(${has_business_col}2&""))="TRUE",'
            f'UPPER(TRIM(${has_business_col}2&""))="1"))'
        )
        business_type_selected = (
            f'LEN(TRIM(${business_type_col}2&""))>0'
        )
        period_blank = f'LEN(TRIM(${business_duration_col}2&""))=0'
        business_duration_validation = DataValidation(
            type="custom",
            formula1=(
                f'OR(AND(NOT({business_yes}),{period_blank}),'
                f'AND({business_yes},NOT({business_type_selected}),{period_blank}),'
                f'AND({business_yes},{business_type_selected},'
                f'ISNUMBER(${business_duration_col}2),'
                f'${business_duration_col}2>=0,${business_duration_col}2<=100))'
            ),
            allow_blank=False,
            showInputMessage=True,
            promptTitle="Select Business Type first",
            prompt=(
                'Select Primary Worker YES, Business Yes and a Type of Business. Enter the period '
                'in years (0 to 100). Decimals such as 0.5 are allowed.'
            ),
            showErrorMessage=True,
            errorStyle="stop",
            errorTitle="Select Business Type first",
            error=(
                'Select Primary Worker YES, Business Yes and a Type of Business first.'
            ),
        )
        worksheet.add_data_validation(business_duration_validation)
        period_range = f"{business_duration_col}2:{business_duration_col}{max_row}"
        business_duration_validation.add(period_range)
        # Validation alone does not flag an untouched cell when Business changes.
        worksheet.conditional_formatting.add(period_range, FormulaRule(
            formula=[f'AND({business_yes},{business_type_selected},{period_blank})'],
            fill=PatternFill(fill_type="solid", fgColor="FFFFC7CE"),
            font=Font(color="FF9C0006"),
        ))
        for column, enabled in (
            (has_business_col, primary_worker_answered),
            (business_type_col, business_yes),
            (business_duration_col, f'AND({business_yes},{business_type_selected})'),
        ):
            target = f"{column}2:{column}{max_row}"
            # Highlight stale answers after changing the worker; never silently erase them.
            worksheet.conditional_formatting.add(target, FormulaRule(
                formula=[f'AND(NOT({enabled}),LEN(TRIM(${column}2&""))>0)'],
                fill=PatternFill(fill_type="solid", fgColor="FFFFC7CE"),
                font=Font(color="FF9C0006"), stopIfTrue=True,
            ))

    def _apply_protection(self, worksheet):
        worksheet.protection.sheet = True
        worksheet.protection.enable()

    def _autosize_columns(self, worksheet):
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                value = cell.value
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 40)

    def _column_letter(self, column_name):
        return get_column_letter(self.columns.index(column_name) + 1)

    def _location_name(self, location, location_type):
        current = location
        while current is not None:
            if getattr(current, "type", None) == location_type:
                return getattr(current, "name", None) or getattr(current, "code", None)
            current = getattr(current, "parent", None)
        return None

    def _micro_catchment_name(self, location):
        """Micro-catchment for the row's location.

        A micro-catchment isn't a level in the Region/TA/GVH/Village chain — it's
        a separate grouping of specific GVHs (and TAs) defined in ``location.MicroCatchment``.
        Prefer a GVH-level link (more specific) and fall back to the TA-level link.
        """
        gvh_location = self._location_ancestor(location, "W")
        if gvh_location is not None:
            name = self._micro_catchment_link_name(gvh_location, "micro_catchments_gvh")
            if name:
                return name

        ta_location = self._location_ancestor(location, "D")
        if ta_location is not None:
            return self._micro_catchment_link_name(ta_location, "micro_catchments_ta")

        return None

    def _location_ancestor(self, location, location_type):
        current = location
        while current is not None:
            if getattr(current, "type", None) == location_type:
                return current
            current = getattr(current, "parent", None)
        return None

    def _micro_catchment_link_name(self, location, related_name):
        related_manager = getattr(location, related_name, None)
        location_id = getattr(location, "id", None)
        if related_manager is None or location_id is None:
            return None
        cache_key = (related_name, location_id)
        if cache_key not in self._micro_catchment_cache:
            link = related_manager.filter(
                validity_to__isnull=True,
                micro_catchment__validity_to__isnull=True,
            ).select_related("micro_catchment").first()
            micro_catchment = link.micro_catchment if link else None
            name = None
            if micro_catchment is not None:
                name = micro_catchment.name or micro_catchment.code
            self._micro_catchment_cache[cache_key] = name
        return self._micro_catchment_cache[cache_key]

    def _hotspot_name(self, location):
        """Hotspot for the row's location.

        A hotspot links to specific villages (``location.HotspotVillage``), so unlike
        the micro-catchment lookup there's no ancestor tier to fall back through —
        just resolve the village-level ancestor's hotspot link, if any.
        """
        village_location = self._location_ancestor(location, "V")
        if village_location is None:
            return None

        related_manager = getattr(village_location, "hotspot_links", None)
        location_id = getattr(village_location, "id", None)
        if related_manager is None or location_id is None:
            return None
        if location_id not in self._hotspot_cache:
            link = related_manager.filter(
                validity_to__isnull=True,
                hotspot__validity_to__isnull=True,
            ).select_related("hotspot").first()
            hotspot = link.hotspot if link else None
            name = None
            if hotspot is not None:
                name = hotspot.name or hotspot.code
            self._hotspot_cache[location_id] = name
        return self._hotspot_cache[location_id]

    def _member_name(self, individual):
        if not individual:
            return None
        first_name = getattr(individual, "first_name", "") or ""
        last_name = getattr(individual, "last_name", "") or ""
        return f"{first_name} {last_name}".strip()

    def _national_id(self, individual):
        if not individual:
            return None
        return (getattr(individual, "json_ext", None) or {}).get("national_id")

    def _marital_status(self, individual):
        if not individual:
            return None
        return (getattr(individual, "json_ext", None) or {}).get("marital_status")

    def _disability(self, individual):
        if not individual:
            return None
        return (getattr(individual, "json_ext", None) or {}).get("disability")

    def _has_business(self, individual):
        if not individual:
            return None
        value = (getattr(individual, "json_ext", None) or {}).get("business_experience")
        if isinstance(value, bool):
            return "Yes" if value else "No"
        if isinstance(value, str):
            normalized = value.strip().upper()
            if normalized == "YES":
                return "Yes"
            if normalized == "NO":
                return "No"
        return None

    def _business_type(self, individual):
        if not individual:
            return None
        return (getattr(individual, "json_ext", None) or {}).get("type_of_business")

    def _business_period(self, individual):
        if not individual:
            return None
        return (getattr(individual, "json_ext", None) or {}).get("business_period")

    def _relationship(self, role):
        if role is None:
            return None
        return str(role).strip() or None

    def _is_head(self, member):
        return str(member.role or "").upper() == "HEAD"

    def _project_name(self, project):
        return getattr(project, "name", None)

    def _project_id(self, project):
        return str(getattr(project, "id", "") or getattr(project, "uuid", "") or "")

    def _project_labels(self):
        name_counts = {}
        for project in self.projects:
            project_name = self._project_name(project)
            if project_name:
                name_counts[project_name] = name_counts.get(project_name, 0) + 1

        labels = {}
        for project in self.projects:
            project_name = self._project_name(project)
            project_id = self._project_id(project)
            if project_name and name_counts.get(project_name, 0) > 1 and project_id:
                labels[id(project)] = f"{project_name} ({project_id})"
            else:
                labels[id(project)] = project_name
        return labels
