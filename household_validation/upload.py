from dataclasses import dataclass, field
import csv
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from io import StringIO

from openpyxl import load_workbook

from household_validation.excel import (
    BUSINESS_DURATION_COLUMN,
    BUSINESS_TYPE_COLUMN,
    EXCEL_COLUMNS,
    HAS_BUSINESS_COLUMN,
    PROJECT_OPTIONS_SHEET,
    PROJECT_OPTIONS_HEADERS,
    _configured_business_type_options,
)
from household_validation.identity import get_household_form_number
from household_validation.wealth import get_household_wealth_quintile
from household_validation.verification import PARTICIPANT_STATUS_COLUMN, HOUSEHOLD_STATUS_COLUMN


VALIDATION_LIST_SHEET = "Validation List"
PROJECT_SELECTION_TYPE_INTENT = "INTENT"
VALIDATION_STATUS_VERIFIED = "VERIFIED"
VALIDATION_STATUS_NOT_VERIFIED = "NOT_VERIFIED"

EDITABLE_UPLOAD_COLUMNS = {
    "national_id",
    "primary_worker",
    "project",
    "validation_notes",
    HAS_BUSINESS_COLUMN,
    BUSINESS_TYPE_COLUMN,
    BUSINESS_DURATION_COLUMN,
}
OPTIONAL_UPLOAD_COLUMNS = {
    PARTICIPANT_STATUS_COLUMN,
    HOUSEHOLD_STATUS_COLUMN,
    "Micro-Catchment",
    "Hotspot",
    "marital_status",
    "disability",
    "pmt_score",
    HAS_BUSINESS_COLUMN,
    BUSINESS_TYPE_COLUMN,
    BUSINESS_DURATION_COLUMN,
}
REQUIRED_UPLOAD_COLUMNS = tuple(
    column for column in EXCEL_COLUMNS if column not in OPTIONAL_UPLOAD_COLUMNS
)
STRUCTURAL_UPLOAD_COLUMNS = tuple(
    column for column in EXCEL_COLUMNS if column not in EDITABLE_UPLOAD_COLUMNS
)

YES_VALUES = {"YES", "Y", "TRUE", "1"}
NO_VALUES = {"NO", "N", "FALSE", "0"}


@dataclass(frozen=True)
class UploadedValidationRow:
    row_number: int
    values: dict
    verified: bool | None
    primary_worker: bool | None
    validation_date: date | None
    project_name: str | None
    project_id: str | None
    notes: str | None
    business_updates: dict = field(default_factory=dict)
    participant_status: str | None = None
    household_status: str | None = None


@dataclass(frozen=True)
class WorkbookParseResult:
    rows: list[UploadedValidationRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    error_row_numbers: frozenset[int] = field(default_factory=frozenset)
    invalid_group_keys: frozenset[str] = field(default_factory=frozenset)
    project_options: dict = field(default_factory=dict)
    total_rows_read: int | None = None

    @property
    def rows_read(self):
        return (
            len(self.rows)
            if self.total_rows_read is None
            else self.total_rows_read
        )


def parse_validation_workbook(file_or_bytes):
    from household_validation.apps import HouseholdValidationConfig

    business_columns_enabled = HouseholdValidationConfig.business_columns_enabled
    workbook = load_workbook(_to_bytes_io(file_or_bytes), data_only=True)
    errors = []
    if VALIDATION_LIST_SHEET not in workbook.sheetnames:
        return WorkbookParseResult(errors=[f"Missing worksheet: {VALIDATION_LIST_SHEET}"])

    worksheet = workbook[VALIDATION_LIST_SHEET]
    try:
        headers = _read_headers(worksheet)
    except ValueError as exc:
        return WorkbookParseResult(errors=[str(exc)])
    missing_columns = [
        column for column in REQUIRED_UPLOAD_COLUMNS if column not in headers
    ]
    if missing_columns:
        return WorkbookParseResult(
            errors=[f"Missing required columns: {', '.join(missing_columns)}"]
        )

    project_options = _read_project_options(workbook)
    rows = []
    error_row_numbers = set()
    invalid_group_keys = set()
    total_rows_read = 0
    for row_number in range(2, worksheet.max_row + 1):
        values = {
            column: (
                _cell_value(
                    worksheet.cell(row=row_number, column=headers[column])
                )
                if column in headers
                else None
            )
            for column in EXCEL_COLUMNS
        }
        if _is_blank_row(values):
            continue
        total_rows_read += 1
        row_errors = _validate_structural_values(row_number, values)
        business_updates, business_errors = _parse_business_values(row_number, values)
        row_errors.extend(business_errors)
        primary_worker = _parse_yes_no(values.get("primary_worker"))
        if business_columns_enabled and primary_worker is not True and any(
            _clean(values.get(column)) is not None
            for column in (HAS_BUSINESS_COLUMN, BUSINESS_TYPE_COLUMN, BUSINESS_DURATION_COLUMN)
        ):
            row_errors.append(
                f"Row {row_number}: business information is only available for the selected "
                "Primary Worker. Clear all three business fields or select Primary Worker YES."
            )
        validation_date = _parse_date(values.get("validation_date"))
        project_label = _clean(values.get("project"))
        project_name = _resolve_project_name(project_label, project_options)
        project_id = _resolve_project_id(
            project_label=project_label,
            workbook_project_id=_clean(values.get("project_id")),
            project_options=project_options,
        )
        if values.get("project_id") and not project_label:
            row_errors.append(f"Row {row_number}: project_id cannot be set without project")
        if values.get("primary_worker") not in (None, "") and _parse_yes_no(
            values.get("primary_worker")
        ) is None:
            row_errors.append(f"Row {row_number}: primary_worker must be YES or NO")
        if values.get("validation_date") and validation_date is None:
            row_errors.append(f"Row {row_number}: validation_date is invalid")
        if project_label and not project_id:
            row_errors.append(f"Row {row_number}: project is not in the project options")
        if row_errors:
            errors.extend(row_errors)
            error_row_numbers.add(row_number)
            group_key = str(values.get("group_uuid") or "").strip()
            if group_key:
                invalid_group_keys.add(group_key)
            continue
        rows.append(
            UploadedValidationRow(
                row_number=row_number,
                values=values,
                # Verification is derived for the whole household by the upload
                # service after projecting all Primary Worker values.
                verified=None,
                primary_worker=primary_worker,
                validation_date=validation_date,
                project_name=project_name,
                project_id=project_id,
                notes=_clean(values.get("validation_notes")),
                business_updates=business_updates,
            )
        )
    return WorkbookParseResult(
        rows=rows,
        errors=errors,
        error_row_numbers=frozenset(error_row_numbers),
        invalid_group_keys=frozenset(invalid_group_keys),
        project_options=project_options,
        total_rows_read=total_rows_read,
    )


def _to_bytes_io(file_or_bytes):
    if isinstance(file_or_bytes, bytes):
        return BytesIO(file_or_bytes)
    if hasattr(file_or_bytes, "read"):
        content = file_or_bytes.read()
        if hasattr(file_or_bytes, "seek"):
            file_or_bytes.seek(0)
        return BytesIO(content)
    return file_or_bytes


def _read_headers(worksheet):
    def normalize(value):
        return " ".join(value.replace("_", " ").split()).casefold()

    canonical = {normalize(column): column for column in EXCEL_COLUMNS}
    canonical.update({
        "business experience": HAS_BUSINESS_COLUMN,
        "has business": HAS_BUSINESS_COLUMN,
        "does member have a business": HAS_BUSINESS_COLUMN,
        "type of business": BUSINESS_TYPE_COLUMN,
        "business type": BUSINESS_TYPE_COLUMN,
        "business period": BUSINESS_DURATION_COLUMN,
    })
    headers = {}
    for column_number in range(1, worksheet.max_column + 1):
        value = _clean(worksheet.cell(row=1, column=column_number).value)
        if value:
            value = canonical.get(normalize(value), value)
            if value in headers:
                raise ValueError(f"Duplicate column: {value}")
            headers[value] = column_number
    return headers


def _parse_business_values(row_number, values):
    updates = {}
    errors = []
    has_business = _clean(values.get(HAS_BUSINESS_COLUMN))
    business_type = _clean(values.get(BUSINESS_TYPE_COLUMN))
    period = _clean(values.get(BUSINESS_DURATION_COLUMN))
    if has_business:
        flag = _parse_yes_no(has_business)
        if flag is None:
            errors.append(f"Row {row_number}: {HAS_BUSINESS_COLUMN} must be YES or NO")
        else:
            updates["business_experience"] = "Yes" if flag else "No"
    if business_type:
        options = {option.casefold(): option for option in _configured_business_type_options()}
        if business_type.casefold() not in options:
            errors.append(f"Row {row_number}: {BUSINESS_TYPE_COLUMN} is not a configured business type")
        else:
            updates["type_of_business"] = options[business_type.casefold()]
    if updates.get("business_experience") == "Yes":
        if not business_type:
            errors.append(
                f"Row {row_number}: {BUSINESS_TYPE_COLUMN} is required when {HAS_BUSINESS_COLUMN} is Yes"
            )
        if not period:
            errors.append(
                f"Row {row_number}: {BUSINESS_DURATION_COLUMN} is required when {HAS_BUSINESS_COLUMN} is Yes"
            )
    if period:
        try:
            number = Decimal(period)
            if not number.is_finite() or not 0 <= number <= 100:
                raise ValueError
            updates["business_period"] = float(number)
        except (InvalidOperation, ValueError):
            errors.append(f"Row {row_number}: {BUSINESS_DURATION_COLUMN} must be between 0 and 100")
        if not business_type:
            errors.append(
                f"Row {row_number}: select {BUSINESS_TYPE_COLUMN} before entering {BUSINESS_DURATION_COLUMN}"
            )
    if updates.get("business_experience") == "No":
        # Changing Yes to No must also remove previously stored business details.
        updates.update(type_of_business=None, business_period=None)
    elif (business_type or period) and updates.get("business_experience") != "Yes":
        errors.append(f"Row {row_number}: select Yes for {HAS_BUSINESS_COLUMN} before entering business details")
    return updates, errors


def _read_project_options(workbook):
    if PROJECT_OPTIONS_SHEET not in workbook.sheetnames:
        return {}
    worksheet = workbook[PROJECT_OPTIONS_SHEET]
    headers = _read_headers(worksheet)
    if not all(header in headers for header in PROJECT_OPTIONS_HEADERS[:2]):
        return {}
    options = {}
    for row_number in range(2, worksheet.max_row + 1):
        project_id = _clean(worksheet.cell(row=row_number, column=headers["project_id"]).value)
        project_name = _clean(worksheet.cell(row=row_number, column=headers["project"]).value)
        project_label = project_name
        if "project_label" in headers:
            project_label = _clean(worksheet.cell(row=row_number, column=headers["project_label"]).value)
        if project_id and project_name:
            options[project_label or project_name] = {
                "id": project_id,
                "name": project_name,
            }
    return options


def _validate_structural_values(row_number, values):
    errors = []
    for column in ("batch_id", "group_uuid", "member_uuid"):
        if not _clean(values.get(column)):
            errors.append(f"Row {row_number}: {column} is required")
    return errors


def _is_blank_row(values):
    return all(value in (None, "") for value in values.values())


def _cell_value(cell):
    value = cell.value
    if isinstance(value, datetime):
        return value.date()
    return value


def _clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _parse_yes_no(value):
    value = _clean(value)
    if value is None:
        return None
    normalized = value.upper()
    if normalized in YES_VALUES:
        return True
    if normalized in NO_VALUES:
        return False
    return None


def _parse_date(value):
    if isinstance(value, date):
        return value
    value = _clean(value)
    if not value:
        return None
    try:
        return date.fromisoformat(value[:10])
    except ValueError:
        return None


def _resolve_project_name(project_label, project_options):
    if not project_label:
        return None
    option = project_options.get(project_label)
    if isinstance(option, dict):
        return option.get("name")
    if option:
        return project_label
    return None


def _resolve_project_id(project_label, workbook_project_id, project_options):
    if not project_label:
        return workbook_project_id
    option = project_options.get(project_label)
    project_id = option.get("id") if isinstance(option, dict) else option
    if workbook_project_id and project_id and workbook_project_id != project_id:
        return None
    return project_id or workbook_project_id


def build_validation_json_ext(uploaded_row, project, upload_date, uploaded_at, user_id):
    validation_date = uploaded_row.validation_date or upload_date
    validation_status = uploaded_row.household_status or (
        VALIDATION_STATUS_VERIFIED
        if uploaded_row.verified is True
        else VALIDATION_STATUS_NOT_VERIFIED
    )
    return {
        "validation_status": validation_status,
        "last_verified_date": validation_date.isoformat(),
        "validation_project_id": str(project.id) if project else uploaded_row.project_id,
        "validation_project_name": project.name if project else uploaded_row.project_name,
        "validation_project_selection_type": PROJECT_SELECTION_TYPE_INTENT,
        "validation_uploaded_at": uploaded_at.isoformat(),
        "validation_uploaded_by_id": str(user_id) if user_id else None,
        "validation_notes": uploaded_row.notes,
    }


def member_structural_errors(uploaded_row, group_individual, group=None):
    errors = []
    row_number = uploaded_row.row_number
    individual = getattr(group_individual, "individual", None)
    individual_json_ext = getattr(individual, "json_ext", None) or {}
    group = group or getattr(group_individual, "group", None)

    expected = {
        "form_number": get_household_form_number(group, individual),
        "member_name": _member_name(individual),
        "member_gender": individual_json_ext.get("gender"),
        "member_dob": _date_value(getattr(individual, "dob", None)),
        "member_age": _age(getattr(individual, "dob", None)),
        "fit_for_work": "YES" if _truthy(individual_json_ext.get("fit_for_work")) else "NO",
        "relationship": _relationship(getattr(group_individual, "role", None)),
        "head": "YES" if str(getattr(group_individual, "role", "")).upper() == "HEAD" else "NO",
        "household_wealth_quintile": get_household_wealth_quintile(group),
    }
    strict_columns = {
        "form_number",
        "relationship",
        "household_wealth_quintile",
    }
    for column, expected_value in expected.items():
        uploaded_value = uploaded_row.values.get(column)
        if expected_value in (None, ""):
            continue
        if column not in strict_columns and uploaded_value in (None, ""):
            continue
        if _normalize(uploaded_value) != _normalize(expected_value):
            errors.append(f"Row {row_number}: {column} does not match the household member")
    return errors


def build_validation_error_report_csv(batch_id, rows):
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "batch_id",
            "row_number",
            "status",
            "form_number",
            "group_uuid",
            "member_uuid",
            "error_message",
        ]
    )
    for row in rows:
        raw_row = row.raw_row or {}
        writer.writerow(
            [
                str(batch_id),
                row.row_number,
                row.status,
                raw_row.get("form_number") or raw_row.get("group_code"),
                raw_row.get("group_uuid"),
                raw_row.get("member_uuid"),
                row.error_message,
            ]
        )
    return output.getvalue()


def _member_name(individual):
    if individual is None:
        return None
    return f"{getattr(individual, 'first_name', '') or ''} {getattr(individual, 'last_name', '') or ''}".strip()


def _relationship(role):
    if role is None:
        return None
    return str(role).strip() or None


def _date_value(value):
    if isinstance(value, date):
        return value.isoformat()
    return value


def _age(dob):
    if not isinstance(dob, date):
        return None
    today = date.today()
    return today.year - dob.year - ((today.month, today.day) < (dob.month, dob.day))


def _truthy(value):
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "y"}
    return bool(value)


def _normalize(value):
    if value is None:
        return None
    if isinstance(value, date):
        value = value.isoformat()
    elif isinstance(value, float) and value.is_integer():
        value = int(value)
    value = str(value).strip().casefold()
    return value or None
